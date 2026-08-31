#!/usr/bin/env python3
"""Build the NYC Public Grocery financial model.

Outputs: public/assets/nyc-public-grocery-model.xlsx

Every assumption on the Assumptions tab is a yellow-highlighted cell. Downstream
tabs pull from those cells by formula, so changing an assumption updates the
full model. Sources are cited in cell comments.
"""
import hashlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

OUTPUT = (
    Path(__file__).resolve().parent.parent.parent
    / "public"
    / "assets"
    / "nyc-public-grocery-model.xlsx"
)

# Freshness stamp: the verification gate compares this against the script on disk,
# so a stale workbook (built from an older script) fails loudly.
SELF_SHA = hashlib.sha256(Path(__file__).resolve().read_bytes()).hexdigest()

# ---- 17A brand + model palette ----
COBALT = "1F1FD6"
CAROLINA = "21A8E0"
CHARCOAL = "3B3B3B"
INPUT_YELLOW = "FFF3CD"
CALC_GRAY = "F5F5F5"
SECTION_BLUE_BG = "E8E8FF"
TOTAL_BG = "D9D9FF"

input_fill = PatternFill("solid", fgColor=INPUT_YELLOW)
calc_fill = PatternFill("solid", fgColor=CALC_GRAY)
section_fill = PatternFill("solid", fgColor=SECTION_BLUE_BG)
total_fill = PatternFill("solid", fgColor=TOTAL_BG)
header_fill = PatternFill("solid", fgColor=COBALT)

thin = Side(style="thin", color="CCCCCC")
bottom = Border(bottom=thin)
top_bottom = Border(top=thin, bottom=thin)

header_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
section_font = Font(name="Calibri", size=11, bold=True, color=COBALT)
label_font = Font(name="Calibri", size=10, color=CHARCOAL)
label_bold = Font(name="Calibri", size=10, color=CHARCOAL, bold=True)
note_font = Font(name="Calibri", size=9, italic=True, color="888888")
total_font = Font(name="Calibri", size=10, bold=True, color=CHARCOAL)

FMT_D = "$#,##0"
FMT_D1 = "$#,##0.00"
FMT_K = '$#,##0,"k"'
FMT_M = '$#,##0.0,,"M"'
FMT_PCT = "0.0%"
FMT_PCT0 = "0%"
FMT_N = "#,##0"
FMT_1 = "0.0"
FMT_SQFT = '#,##0" sqft"'


def put_header(ws, row, text, span=8):
    ws.cell(row=row, column=1, value=text).font = header_font
    ws.cell(row=row, column=1).fill = header_fill
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 26
    for c in range(2, span + 1):
        ws.cell(row=row, column=c).fill = header_fill


def put_section(ws, row, text, span=8):
    c = ws.cell(row=row, column=1, value=text)
    c.font = section_font
    c.fill = section_fill
    for col in range(2, span + 1):
        ws.cell(row=row, column=col).fill = section_fill


def put_input(ws, row, col, value, fmt=None, comment=None):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = input_fill
    c.font = label_bold
    if fmt:
        c.number_format = fmt
    if comment:
        c.comment = Comment(comment, "17A")
    return c


def put_calc(ws, row, col, value, fmt=None, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = calc_fill
    c.font = total_font if bold else label_font
    if fmt:
        c.number_format = fmt
    return c


def put_label(ws, row, col, text, bold=False, note=False):
    c = ws.cell(row=row, column=col, value=text)
    if note:
        c.font = note_font
    elif bold:
        c.font = label_bold
    else:
        c.font = label_font
    return c


def put_total(ws, row, col, value, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = total_fill
    c.font = total_font
    c.border = top_bottom
    if fmt:
        c.number_format = fmt
    return c


wb = Workbook()

# =============================================================================
# 1. README
# =============================================================================
ws = wb.active
ws.title = "README"
ws.column_dimensions["A"].width = 95

put_header(ws, 1, "NYC Public Grocery Plan — Financial Model", span=2)

readme_blocks = [
    ("What this is",
     "A pro-forma P&L for Mayor Mamdani's plan to build five city-owned grocery "
     "stores, stress-tested against the same budget spent on six alternative "
     "food-access interventions. Built to make the assumptions legible so any "
     "reader can disagree with a number, change it, and see where the "
     "conclusion moves."),
    ("How to use",
     "Inputs live on the Assumptions tab (yellow fill). Downstream tabs are "
     "formula-driven from those inputs wherever feasible; the handful of "
     "scenario literals that remain are flagged in the notes where they occur. "
     "Each assumption has a cell-comment with its source."),
    ("Tabs",
     "• Assumptions — all inputs, grouped and sourced.\n"
     "• Per-Store P&L — a single 9,000 sqft store over 10 years, with ramp.\n"
     "• Five-Store Rollup — all five stores phased to the announced timeline.\n"
     "• Sensitivity — steady-state operating result by sales/sqft and discount.\n"
     "• July 2026 Update — the announced 30% core-basket discount: blended\n"
     "  scenarios, per-store losses, leverage in today's dollars, horizon cases,\n"
     "  and the same money spent other ways.\n"
     "• Alternatives — six other interventions ranked against the plan at both\n"
     "  of its centrals (the April piece's published reach × spend frame; the\n"
     "  July 2026 Update tab carries the today's-dollars metric).\n"
     "• Sources — full citation list."),
    ("Key sources",
     "• NYCEDC, N.Y.C. Groceries Operator(s) RFP (July 2026) — footprints,\n"
     "  affordability-payment structure, basket definition, pricing rules.\n"
     "• N.Y.C. Groceries vision plan, 'A Recipe for Affordability' (July 27, 2026).\n"
     "• FMI 2024/2025 Food Industry Facts — sales/sqft, margin, store sizes.\n"
     "• USDA Thrifty Food Plan, 2025 monthly cost of food reports.\n"
     "• NYS OTDA — SNAP caseloads, NYC (2025).\n"
     "• NYC Comptroller, Fiscal Note 4-2024 (The FRESH Program, Oct 2024).\n"
     "• NY Department of Labor — 2026 prevailing wage schedules.\n"
     "• Company filings via stockanalysis.com — Kroger, Costco net margins.\n"
     "• Mayor's Office, La Marqueta announcement (April 14, 2026).\n"
     "• Mayor's Office, 30% core-basket discount announcement (July 27, 2026).\n"
     "• Kansas City Sun Fresh closure (KCUR, 2025); Baldwin FL closure; "
     "Chicago municipal grocery shelved (Supermarket News, 2024)."),
    ("Build info",
     "Generated by gizmos/nyc-public-grocery-math/build_model.py in the "
     "gizmo-warehouse repo. "
     "Author: Joe Eichenbaum, 17A. Questions: joe@group17a.com."),
    ("Build fingerprint",
     f"build fingerprint (sha256 of build_model.py at build time): {SELF_SHA}. "
     "Regenerate the workbook after any change to the script — the verification "
     "gate compares this stamp against the script on disk."),
]

r = 3
for heading, body in readme_blocks:
    c = ws.cell(row=r, column=1, value=heading)
    c.font = section_font
    c.fill = section_fill
    r += 1
    c = ws.cell(row=r, column=1, value=body)
    c.font = label_font
    c.alignment = Alignment(wrap_text=True, vertical="top")
    # +1 wrap line of headroom: LibreOffice/Excel wrap slightly wider than the
    # 85-chars-per-line estimate, which clipped the last line of short blocks
    ws.row_dimensions[r].height = 15 * (body.count("\n") + max(1, len(body) // 85) + 1)
    r += 2


# =============================================================================
# 2. ASSUMPTIONS
# =============================================================================
ws = wb.create_sheet("Assumptions")
ws.column_dimensions["A"].width = 44
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 70

put_header(ws, 1, "Assumptions — all inputs in yellow", span=3)
put_label(ws, 2, 1, "Change any yellow cell and downstream tabs update.", note=True)

# Track named cells for easy reference
A = {}  # name -> cell address like "Assumptions!$B$5"

def add_named(key, row):
    A[key] = f"Assumptions!$B${row}"

r = 4

# ---- STORE PHYSICAL ----
put_section(ws, r, "Store physical", span=3); r += 1
put_label(ws, r, 1, "Selling square footage per store", bold=True)
put_input(ws, r, 2, 9000, fmt=FMT_SQFT,
          comment="La Marqueta flagship footprint per Mayor's announcement (April 14, 2026). "
                  "The May 18, 2026 release and EDC portal carried 20,000 sqft for The "
                  "Peninsula and 10,000 sqft minimums for the rest; the July 27, 2026 "
                  "RFP respecifies the Peninsula at ~15,000 sqft and instructs bidders "
                  "to assume 15,000 sqft selling for the unnamed sites — see the July "
                  "2026 Update tab (13,800 sqft average).")
put_label(ws, r, 3, "FMI industry average full-format store: ~42,000 sqft.", note=True); add_named("sqft", r); r += 1
put_label(ws, r, 1, "Number of stores (plan)", bold=True)
put_input(ws, r, 2, 5, fmt=FMT_N, comment="One per borough per Mayor Mamdani's plan.")
put_label(ws, r, 3, "Phased 2027–2029.", note=True); add_named("n_stores", r); r += 2

# ---- REVENUE ----
put_section(ws, r, "Revenue assumptions", span=3); r += 1
put_label(ws, r, 1, "Steady-state revenue per sqft ($/yr)", bold=True)
put_input(ws, r, 2, 500, fmt=FMT_D,
          comment="FMI 2024 industry average was $965/sqft ($18.55/sqft/week); "
                  "FMI 2025 puts it at $1,019 — see the July 2026 Update tab. "
                  "Those figures reflect full-format 42k sqft supermarkets. "
                  "Small-format stores in low-income neighborhoods without "
                  "brand pull (Trader Joe's, WF Daily Shop) typically realize "
                  "$400–600/sqft in their first years. Central case: $500.")
put_label(ws, r, 3, "Industry avg: $1,019/sqft (FMI 2025; $965 in FMI 2024). Small-format realistic: $400–600.", note=True); add_named("rev_sqft", r); r += 1
put_label(ws, r, 1, "Year 1 ramp (% of steady-state)", bold=True)
put_input(ws, r, 2, 0.60, fmt=FMT_PCT, comment="New stores typically open below stabilized revenue; 60% in Y1 is common.")
add_named("ramp_y1", r); r += 1
put_label(ws, r, 1, "Year 2 ramp (% of steady-state)", bold=True)
put_input(ws, r, 2, 0.80, fmt=FMT_PCT); add_named("ramp_y2", r); r += 1
put_label(ws, r, 1, "Discount to shoppers vs. market price", bold=True)
put_input(ws, r, 2, 0.10, fmt=FMT_PCT,
          comment="When this model was built (April 2026), the administration had "
                  "named no percentage — campaign and April 2026 materials promised "
                  "only 'wholesale prices' / passing on savings. 10% was this model's "
                  "central-case assumption. On July 27, 2026 the administration "
                  "announced 30% off a core basket (~15% off the average total bill, "
                  "per City Hall) — see the July 2026 Update tab. This cell stays at "
                  "the April central case; note that after the August 2026 re-audit "
                  "(school meals, FRESH, SNAP count corrected; Alternatives tab "
                  "rebuilt) the April piece's tables match this workbook only where "
                  "that piece's banner says so. "
                  "Discount applies to realized revenue line (price cut, not cost cut).")
put_label(ws, r, 3, "No official % existed pre-July 2026. Announced 2026-07-27: 30% core basket — see July 2026 Update tab.", note=True); add_named("discount", r); r += 2

# ---- COST STRUCTURE ----
put_section(ws, r, "Cost structure (% of GROSS revenue — i.e., at market prices)", span=3); r += 1
put_label(ws, r, 3, "Costs are per-unit. Applying a shopper discount reduces revenue but not cost of goods. All cost %s are anchored to gross (pre-discount) revenue so the discount flows through to margin correctly.", note=True); r += 1
put_label(ws, r, 1, "COGS (% of gross revenue)", bold=True)
put_input(ws, r, 2, 0.72, fmt=FMT_PCT,
          comment="FMI 2024: industry-average COGS ~72–73% at market prices. "
                  "A municipal store without warehouse-club purchasing scale "
                  "probably runs higher. 72% is charitable — it matches "
                  "large-chain average despite small-format disadvantage.")
add_named("cogs_pct", r); r += 1
put_label(ws, r, 1, "Labor (% of gross revenue)", bold=True)
put_input(ws, r, 2, 0.12, fmt=FMT_PCT,
          comment="FMI industry avg labor ~11% of revenue. NYC prevailing wage "
                  "adds ~15–25% to the base — ~12% of gross revenue is a charitable "
                  "central case. Sensitivity range: 10–14%.")
add_named("labor_pct", r); r += 1
put_label(ws, r, 1, "Other opex (utilities, shrink, insurance, tech) — ex rent", bold=True)
put_input(ws, r, 2, 0.09, fmt=FMT_PCT,
          comment="FMI 2024 — utilities ~1.2%, shrink ~1.6%, insurance/tech/misc ~6%. "
                  "Rent excluded (city-owned). ~9% ex-rent.")
add_named("opex_pct", r); r += 1
put_label(ws, r, 1, "Private operator fee (annual, per store)", bold=True)
put_input(ws, r, 2, 400000, fmt=FMT_D,
          comment="Plan assumes no profit margin but private operator runs day-to-day. "
                  "Operator fee ~$400k/yr covers management overhead. Plausible range "
                  "$200–600k. July 2026 update: the operator RFP structures compensation "
                  "as annual 'affordability payments' compensating the discount loss, "
                  "sized by each bidder's own 'requests for subsidy' (a scored selection "
                  "factor), on top of free rent, no property tax, and city-funded "
                  "buildout. This model's loss-based subsidy is therefore CONSERVATIVE — "
                  "full discount reimbursement would cost more (see the July 2026 Update "
                  "tab's reimbursement-structure row and IMPLEMENTATION_LEDGER_2026-08.md).")
add_named("op_fee", r); r += 2

# ---- MEMO: RENT ----
put_section(ws, r, "Rent savings (memo only — city-owned land)", span=3); r += 1
put_label(ws, r, 1, "Market-rate rent comp ($/sqft/yr)", bold=True)
put_input(ws, r, 2, 50, fmt=FMT_D,
          comment="Retail rent in East Harlem / South Bronx: $40–60/sqft. $50 central case.")
add_named("rent_comp", r); r += 1
put_label(ws, r, 1, "Implied annual rent avoided (per store)", bold=True)
put_calc(ws, r, 2, f"={A['rent_comp']}*{A['sqft']}", fmt=FMT_D, bold=True)
add_named("rent_memo", r); r += 2

# ---- CAPEX ----
put_section(ws, r, "Capital cost", span=3); r += 1
put_label(ws, r, 1, "Capex — La Marqueta flagship", bold=True)
put_input(ws, r, 2, 30000000, fmt=FMT_D,
          comment="Per Mayor's April 14, 2026 announcement. $30M for 9,000 sqft = "
                  "$3,333/sqft. Benchmarks (corrected Aug 2026): grocery fit-out of "
                  "existing space $117–211/sqft; ground-up national new-build "
                  "~$215–320/sqft; NYC all-in (hard + soft costs) plausibly "
                  "$600–1,200/sqft. Against like-for-like NYC builds, La Marqueta "
                  "runs roughly 3–6× — not the 6–22× an out-of-date $150–500 "
                  "benchmark implied.")
add_named("capex_flagship", r); r += 1
put_label(ws, r, 1, "Capex — other four stores (avg)", bold=True)
put_input(ws, r, 2, 10000000, fmt=FMT_D,
          comment="Budget split implies $10M per store ($70M total − $30M flagship = "
                  "$40M ÷ 4). These four will convert existing retail spaces rather "
                  "than new construction (per Mayor's Office, April 2026), so actual "
                  "per-store cost is probably lower — covering lease buyouts, tenant "
                  "fit-outs, refrigeration, fixtures, inventory, working capital. "
                  "Exact per-store figures not yet disclosed; adjust this cell as "
                  "information becomes available.")
add_named("capex_other", r); r += 1
put_label(ws, r, 1, "Implied capex per sqft — flagship", bold=True)
put_calc(ws, r, 2, f"={A['capex_flagship']}/{A['sqft']}", fmt=FMT_D, bold=True)
put_label(ws, r, 3, "Benchmarks: fit-out $117–211/sqft; ground-up national ~$215–320; NYC all-in plausibly $600–1,200.", note=True); r += 2

# ---- DISCOUNT RATE / NPV ----
put_section(ws, r, "Discounting", span=3); r += 1
put_label(ws, r, 1, "Municipal discount rate (for NPV)", bold=True)
put_input(ws, r, 2, 0.04, fmt=FMT_PCT,
          comment="NYC municipal bond yields ~3.5–4.5% (2026). 4% central case.")
add_named("disc_rate", r); r += 2

# ---- ALTERNATIVES ----
put_section(ws, r, "Alternatives — unit costs & reach", span=3); r += 1
put_label(ws, r, 1, "School meals — full public cost per new meal", bold=True)
put_input(ws, r, 2, 4.91, fmt=FMT_D1,
          comment="CACFP At-Risk afterschool supper rate, SY2025-26 — the closest "
                  "funded stream for weekend meals at area-eligible sites, used here "
                  "as the full-public-cost analog. Federal dollars are public money, "
                  "so new meals are priced at the full rate, not the state increment. "
                  "MEMO: the NY UFSM state top-up alone runs ~$0.35 per meal "
                  "opportunity (~$0.86 per meal served on the FY26 $340M "
                  "appropriation) — that was this cell's pre-August-2026 value, and "
                  "it counted only the state's slice of a mostly-federal bill.")
put_label(ws, r, 3, "Full public cost. The state UFSM top-up (~$0.35–0.86/meal) is memo only, no longer used in any row.", note=True)
add_named("meal_cost", r); r += 1
put_label(ws, r, 1, "Meals per kid per year (weekend expansion)", bold=True)
put_input(ws, r, 2, 160, fmt=FMT_N,
          comment="2 meals/day × 2 weekend days × 40 school-year weekends = 160. "
                  "Summer dropped (Aug 2026): NYC summer meals are already universal "
                  "via SFSP, plus SUN Bucks EBT. CACFP At-Risk actually funds one "
                  "meal + one snack per day — this row prices meals at the CACFP "
                  "rate as a cost analog, not a claim that one stream covers all of "
                  "them.")
add_named("meals_per_kid", r); r += 1
put_label(ws, r, 1, "SNAP top-up per household per month", bold=True)
put_input(ws, r, 2, 20, fmt=FMT_D,
          comment="Modest monthly supplement reduces benefits-cliff impact for families with kids.")
add_named("snap_topup", r); r += 1
put_label(ws, r, 1, "Warehouse-club annual membership (Costco Gold Star)", bold=True)
put_input(ws, r, 2, 65, fmt=FMT_D, comment="Costco Gold Star 2026.")
add_named("warehouse_fee", r); r += 1
put_label(ws, r, 1, "Avg household grocery savings from warehouse club (memo — not wired)", bold=True)
put_input(ws, r, 2, 1000, fmt=FMT_D,
          comment="The documented evidence is percentage wedges, not dollars: club "
                  "prices ran 20–35% below supermarket chains on a unit-price-matched "
                  "basket (Consumers' Checkbook 2022), and Consumer Reports (2025 "
                  "pricing) found Costco/BJ's ~21% below even Walmart. The dollar "
                  "figure here is OUR conversion of those wedges at $4–6k/yr of "
                  "food-at-home spend; $800–1,200 is the plausible range.")
add_named("warehouse_savings", r); r += 1
put_label(ws, r, 1, "NYC SNAP households (total)", bold=True)
put_input(ws, r, 2, 1069000, fmt=FMT_N,
          comment="OTDA SNAP caseloads dataset dq6j-8u8z (data.ny.gov), NYC "
                  "(HRA-administered district), 2025 monthly average ≈ 1,069,000 "
                  "households; recipients ≈ 1.77M, i.e. ~1.66 persons per SNAP "
                  "household. Corrected Aug 2026 from 692,000, a survey-scale "
                  "person/household mix-up misattributed to HRA.")
add_named("snap_households", r); r += 1
put_label(ws, r, 1, "NYCHA students (weekend school-meals row)", bold=True)
put_input(ws, r, 2, 62000, fmt=FMT_N,
          comment="NYCHA conventional public housing ~298k residents (341k incl. "
                  "PACT), ~20.7% school-age → ~62,000 students (NYCHA Fact Sheet "
                  "2025). Re-based August 2026 from the April model's 110,000.")
put_label(ws, r, 3, "Drives the Alternatives tab's weekend school-meals row (priced at full public cost).", note=True)
add_named("nycha_students", r); r += 1
put_label(ws, r, 1, "Bodegas in NYC", bold=True)
put_input(ws, r, 2, 10000, fmt=FMT_N,
          comment="N.Y.C. Groceries vision plan, p. 14: NYC is 'home to more than "
                  "1,100 grocery stores and over 10,000 bodegas.'")
add_named("bodegas", r); r += 1
put_label(ws, r, 1, "FRESH program public cost per new store", bold=True)
put_input(ws, r, 2, 1080000, fmt=FMT_D,
          comment="NYC Comptroller, Fiscal Note 4-2024 (Oct 2024): $29.2M in tax "
                  "expenditures through FY2023 ÷ 27 tax-subsidized stores ≈ "
                  "$1.08M/store (realized basis; 25-yr abatement commitments run "
                  "higher). The '30 stores since 2009' and '$177M private capital' "
                  "figures are NYCEDC program-page numbers, not the Comptroller's — "
                  "the pre-Aug-2026 '$45M ÷ 30 = $1.5M' derivation appears in "
                  "neither source.")
add_named("fresh_per_store", r); r += 2

# ---- COMPARISON ENVELOPE ----
put_section(ws, r, "Comparison budget envelope (5-store plan footprint)", span=3); r += 1
put_label(ws, r, 1, "Total capex (5 stores)", bold=True)
put_calc(ws, r, 2, f"={A['capex_flagship']}+4*{A['capex_other']}", fmt=FMT_D, bold=True)
add_named("total_capex", r); r += 1
put_label(ws, r, 1, "Annual operating subsidy (midpoint estimate)", bold=True)
put_input(ws, r, 2, 3000000, fmt=FMT_D,
          comment="Central-case estimate: $600k/store/yr × 5 stores. Range: $1.5–$5M/yr "
                  "depending on realized sales/sqft and discount.")
add_named("annual_subsidy", r); r += 1
put_label(ws, r, 1, "10-yr total budget (capex + 10×opex)", bold=True)
put_calc(ws, r, 2, f"={A['total_capex']}+10*{A['annual_subsidy']}", fmt=FMT_D, bold=True)
add_named("total_budget", r); r += 2

# ---- ALTERNATIVES REACH & PACE (added Aug 2026; appended so rows above keep
# their addresses — downstream tabs and the gate reference them by cell) ----
put_section(ws, r, "Alternatives — reach & pace inputs (added Aug 2026)", span=3); r += 1
put_label(ws, r, 1, "NYC total households", bold=True)
put_input(ws, r, 2, 3334088, fmt=FMT_N,
          comment="Census QuickFacts / ACS 2020–2024 5-yr: 3,334,088 NYC households. "
                  "The deduped ceiling for any walking-distance reach claim.")
add_named("nyc_households", r); r += 1
put_label(ws, r, 1, "Warehouse-club utilization (author assumption)", bold=True)
put_input(ws, r, 2, 0.70, fmt=FMT_PCT0,
          comment="AUTHOR ASSUMPTION, not a documented rate. Documented analogs run "
                  "wide: subsidized/discounted membership take-up 7–46%; necessity "
                  "program take-up 75–88%. 70% is on the generous end; the row "
                  "scales linearly if you disagree. Single source of truth — the "
                  "July 2026 Update and Alternatives tabs both read this cell.")
add_named("warehouse_util", r); r += 1
put_label(ws, r, 1, "FRESH pace — new stores per decade (memo — not wired to any calculation)", bold=True)
put_input(ws, r, 2, 50, fmt=FMT_N,
          comment="Scenario context only; no formula reads this cell. Historic pace "
                  "~2/yr (27 tax-subsidized stores since 2009, Comptroller FN "
                  "4-2024). The July tab's same-money FRESH line runs ~15/yr — its "
                  "own note carries the ceiling-not-forecast caveat.")
add_named("fresh_pace", r); r += 1
put_label(ws, r, 1, "FRESH catchment per store (HH within walking distance; memo — not wired)", bold=True)
put_input(ws, r, 2, 15000, fmt=FMT_N,
          comment="Catchment arithmetic double-counts overlapping catchments. "
                  "NYC households beyond a half-mile walk to a supermarket per "
                  "USDA FARA: ~143k all incomes, of which ~32k low-income.")
add_named("fresh_hh_per_store", r); r += 1
put_label(ws, r, 1, "Health Bucks expansion participants (scenario)", bold=True)
put_input(ws, r, 2, 200000, fmt=FMT_N,
          comment="Author scenario ≈25× the documented program's distribution scale. "
                  "Leverage is 1.0x by construction (direct $-for-$ match), so the "
                  "ranking is scale-invariant.")
add_named("hb_reach", r); r += 1
put_label(ws, r, 1, "Share of SNAP households with kids", bold=True)
put_input(ws, r, 2, 0.40, fmt=FMT_PCT0,
          comment="USDA FNS SNAP household characteristics; drives the supplement "
                  "row's targeting.")
add_named("snap_kids_share", r); r += 1


# =============================================================================
# 3. PER-STORE 10-YEAR P&L
# =============================================================================
ws = wb.create_sheet("Per-Store P&L")
ws.column_dimensions["A"].width = 42
for col in range(2, 13):
    ws.column_dimensions[get_column_letter(col)].width = 13

put_header(ws, 1, "Per-store 10-year P&L (single 9,000 sqft store)", span=12)
put_label(ws, 2, 1, "All lines pulled from Assumptions. Year 0 = pre-opening capex.", note=True)

# Year row
r = 4
put_label(ws, r, 1, "Year", bold=True)
for i in range(11):
    c = ws.cell(row=r, column=2 + i, value=i)
    c.font = label_bold
    c.alignment = Alignment(horizontal="right")
r += 1
put_label(ws, r, 1, "Ramp factor", bold=True)
for i in range(11):
    if i == 0:
        put_calc(ws, r, 2 + i, 0, fmt=FMT_PCT)
    elif i == 1:
        put_calc(ws, r, 2 + i, f"={A['ramp_y1']}", fmt=FMT_PCT)
    elif i == 2:
        put_calc(ws, r, 2 + i, f"={A['ramp_y2']}", fmt=FMT_PCT)
    else:
        put_calc(ws, r, 2 + i, 1.0, fmt=FMT_PCT)
ramp_row = r
r += 2

# Revenue
put_label(ws, r, 1, "Gross revenue (pre-discount)", bold=True)
for i in range(11):
    col = 2 + i
    if i == 0:
        put_calc(ws, r, col, 0, fmt=FMT_D)
    else:
        f = f"={A['sqft']}*{A['rev_sqft']}*{get_column_letter(col)}{ramp_row}"
        put_calc(ws, r, col, f, fmt=FMT_D)
gross_rev_row = r
r += 1
put_label(ws, r, 1, "Less: shopper discount", bold=True)
for i in range(11):
    col = 2 + i
    put_calc(ws, r, col, f"=-{get_column_letter(col)}{gross_rev_row}*{A['discount']}", fmt=FMT_D)
discount_row = r
r += 1
put_label(ws, r, 1, "Net revenue", bold=True)
for i in range(11):
    col = 2 + i
    put_total(ws, r, col,
              f"={get_column_letter(col)}{gross_rev_row}+{get_column_letter(col)}{discount_row}",
              fmt=FMT_D)
rev_row = r
r += 2

# Costs (computed against GROSS revenue — per-unit costs unaffected by price discount)
put_label(ws, r, 1, "COGS (% of gross)", bold=True)
for i in range(11):
    col = 2 + i
    put_calc(ws, r, col, f"=-{get_column_letter(col)}{gross_rev_row}*{A['cogs_pct']}", fmt=FMT_D)
cogs_row = r
r += 1
put_label(ws, r, 1, "Gross margin", bold=True)
for i in range(11):
    col = 2 + i
    put_calc(ws, r, col,
             f"={get_column_letter(col)}{rev_row}+{get_column_letter(col)}{cogs_row}",
             fmt=FMT_D, bold=True)
gm_row = r
r += 2

put_label(ws, r, 1, "Labor (% of gross)", bold=True)
for i in range(11):
    col = 2 + i
    put_calc(ws, r, col, f"=-{get_column_letter(col)}{gross_rev_row}*{A['labor_pct']}", fmt=FMT_D)
labor_row = r
r += 1
put_label(ws, r, 1, "Other opex (utilities, shrink, insurance) — % of gross", bold=True)
for i in range(11):
    col = 2 + i
    put_calc(ws, r, col, f"=-{get_column_letter(col)}{gross_rev_row}*{A['opex_pct']}", fmt=FMT_D)
oth_row = r
r += 1
put_label(ws, r, 1, "Private operator fee", bold=True)
for i in range(11):
    col = 2 + i
    if i == 0:
        put_calc(ws, r, col, 0, fmt=FMT_D)
    else:
        put_calc(ws, r, col, f"=-{A['op_fee']}", fmt=FMT_D)
fee_row = r
r += 2

# Operating result
put_label(ws, r, 1, "Operating result", bold=True)
for i in range(11):
    col = 2 + i
    f = (f"={get_column_letter(col)}{gm_row}+{get_column_letter(col)}{labor_row}+"
         f"{get_column_letter(col)}{oth_row}+{get_column_letter(col)}{fee_row}")
    put_total(ws, r, col, f, fmt=FMT_D)
opres_row = r
r += 2

# Capex
put_label(ws, r, 1, "Capex (flagship)", bold=True)
for i in range(11):
    col = 2 + i
    if i == 0:
        put_calc(ws, r, col, f"=-{A['capex_flagship']}", fmt=FMT_D)
    else:
        put_calc(ws, r, col, 0, fmt=FMT_D)
capex_row = r
r += 1

# Net cash flow
put_label(ws, r, 1, "Net cash flow", bold=True)
for i in range(11):
    col = 2 + i
    put_total(ws, r, col,
              f"={get_column_letter(col)}{opres_row}+{get_column_letter(col)}{capex_row}",
              fmt=FMT_D)
ncf_row = r
r += 2

# Summary metrics
put_label(ws, r, 1, "10-year undiscounted cash flow", bold=True)
ws.cell(row=r, column=2, value=f"=SUM({get_column_letter(2)}{ncf_row}:{get_column_letter(12)}{ncf_row})")
ws.cell(row=r, column=2).font = total_font
ws.cell(row=r, column=2).fill = total_fill
ws.cell(row=r, column=2).number_format = FMT_D
r += 1

put_label(ws, r, 1, "10-year NPV @ municipal discount rate", bold=True)
ws.cell(row=r, column=2,
        value=f"=NPV({A['disc_rate']},{get_column_letter(3)}{ncf_row}:{get_column_letter(12)}{ncf_row})+{get_column_letter(2)}{ncf_row}")
ws.cell(row=r, column=2).font = total_font
ws.cell(row=r, column=2).fill = total_fill
ws.cell(row=r, column=2).number_format = FMT_D
r += 1

put_label(ws, r, 1, "Memo: rent avoided (not in P&L)", bold=True)
for i in range(11):
    col = 2 + i
    if i == 0:
        put_calc(ws, r, col, 0, fmt=FMT_D)
    else:
        put_calc(ws, r, col, f"={A['rent_memo']}", fmt=FMT_D)
r += 1
put_label(ws, r, 1, "Rent savings as % of net revenue (steady state)", bold=True)
put_calc(ws, r, 2, f"={A['rent_memo']}/({A['sqft']}*{A['rev_sqft']}*(1-{A['discount']}))",
         fmt=FMT_PCT, bold=True)


# =============================================================================
# 4. FIVE-STORE ROLLUP
# =============================================================================
ws = wb.create_sheet("Five-Store Rollup")
ws.column_dimensions["A"].width = 36
for col in range(2, 13):
    ws.column_dimensions[get_column_letter(col)].width = 13

put_header(ws, 1, "Five-store rollup — phased to announced timeline", span=12)
put_label(ws, 2, 1, "Assumes The Peninsula (Hunts Point) opens late 2027 (Y1); stores 2–5, including La Marqueta, open by 2029 (Y3). Stores reach steady state by Y3 of operation.", note=True)

# Timeline: which stores are operating in which year
# Year:                2027  2028  2029  2030  2031  2032  2033  2034  2035  2036  2037
# (Model years):         1     2     3     4     5     6     7     8     9    10    11 (unused)
# Store 1 op year:      Y1    Y2    Y3    Y4    Y5    Y6    Y7    Y8    Y9    Y10   Y11
# Store 2 op year:            Y1    Y2    Y3    ...
# Store 3 op year:            Y1    Y2    Y3    ...
# Store 4 op year:                  Y1    Y2    Y3    ...
# Store 5 op year:                  Y1    Y2    Y3    ...

# Simplified: store operating year per calendar year (0 = not operating)
# Columns: Year 0 (2026, pre-capex) through Year 10 (2036)
# Actually let's use Y0 (2026 planning, flagship capex) ... Y10 (2036)
# Opening schedule:
#   Store 1 opens Y1, so operates Y1-Y10 as op-years 1..10
#   Stores 2 and 3 open Y2, operate op-years 1..9 in model years Y2..Y10
#   Stores 4 and 5 open Y3, operate op-years 1..8 in model years Y3..Y10
#   Each store incurs capex the year before opening.

# Build the rollup:
r = 4
put_label(ws, r, 1, "Year", bold=True)
for i in range(11):
    c = ws.cell(row=r, column=2 + i, value=f"Y{i}")
    c.font = label_bold
    c.alignment = Alignment(horizontal="right")
r += 1
put_label(ws, r, 1, "Calendar year", bold=True)
for i in range(11):
    c = ws.cell(row=r, column=2 + i, value=2026 + i)
    c.font = label_font
    c.alignment = Alignment(horizontal="right")
r += 1
put_label(ws, r, 1, "Stores operating", bold=True)
stores_schedule = [0, 1, 3, 5, 5, 5, 5, 5, 5, 5, 5]
for i, n in enumerate(stores_schedule):
    put_calc(ws, r, 2 + i, n, fmt=FMT_N)
stores_row = r
r += 2

# Capex schedule: Y0 = flagship, Y1 = stores 2-3, Y2 = stores 4-5
capex_schedule_f = [
    f"=-{A['capex_flagship']}",  # Y0
    f"=-2*{A['capex_other']}",    # Y1
    f"=-2*{A['capex_other']}",    # Y2
] + ["=0"] * 8

put_label(ws, r, 1, "Capex", bold=True)
for i, f in enumerate(capex_schedule_f):
    put_calc(ws, r, 2 + i, f, fmt=FMT_D)
capex_row = r
r += 1

# For each store, compute its revenue and opex based on its operating year in each calendar year.
# Simplification: treat all 5 stores as identical to the flagship 9,000 sqft, $500/sqft.
# Operating year for each store in each year:
# Store index (1-5): opens Y1, Y2, Y2, Y3, Y3
opens = [1, 2, 2, 3, 3]

# For the rollup, let's compute: total net revenue per calendar year = sum across stores of
# each store's revenue given its operating year. Operating year 1 = ramp_y1, op year 2 = ramp_y2, op year 3+ = 1.0
# Instead of modeling each store separately, express total as:
# revenue(Y) = sqft * rev_sqft * sum_over_stores(ramp(operating_year_of_store_in_Y))
# We'll compute the aggregate ramp factor per calendar year.

# Aggregate ramp per calendar year:
def store_op_year(store_open_year, calendar_year):
    if calendar_year < store_open_year:
        return 0
    return calendar_year - store_open_year + 1

def ramp_formula_for_op_year(op_year):
    if op_year == 0:
        return "0"
    elif op_year == 1:
        return A['ramp_y1']
    elif op_year == 2:
        return A['ramp_y2']
    else:
        return "1"

put_label(ws, r, 1, "Aggregate ramp-adjusted store-years", bold=True)
for i in range(11):
    col = 2 + i
    parts = []
    for open_year in opens:
        oy = store_op_year(open_year, i)
        parts.append(ramp_formula_for_op_year(oy))
    f = "=" + "+".join(parts)
    put_calc(ws, r, col, f, fmt=FMT_1)
agg_row = r
r += 1

put_label(ws, r, 1, "Gross revenue (all stores)", bold=True)
for i in range(11):
    col = 2 + i
    f = f"={A['sqft']}*{A['rev_sqft']}*{get_column_letter(col)}{agg_row}"
    put_calc(ws, r, col, f, fmt=FMT_D)
gross_row = r
r += 1

put_label(ws, r, 1, "Net revenue (after discount)", bold=True)
for i in range(11):
    col = 2 + i
    f = f"={get_column_letter(col)}{gross_row}*(1-{A['discount']})"
    put_calc(ws, r, col, f, fmt=FMT_D)
net_rev_row = r
r += 2

put_label(ws, r, 1, "COGS (% of gross)", bold=True)
for i in range(11):
    col = 2 + i
    put_calc(ws, r, col, f"=-{get_column_letter(col)}{gross_row}*{A['cogs_pct']}", fmt=FMT_D)
cogs_r = r
r += 1
put_label(ws, r, 1, "Labor (% of gross)", bold=True)
for i in range(11):
    col = 2 + i
    put_calc(ws, r, col, f"=-{get_column_letter(col)}{gross_row}*{A['labor_pct']}", fmt=FMT_D)
labor_r = r
r += 1
put_label(ws, r, 1, "Other opex (% of gross)", bold=True)
for i in range(11):
    col = 2 + i
    put_calc(ws, r, col, f"=-{get_column_letter(col)}{gross_row}*{A['opex_pct']}", fmt=FMT_D)
oth_r = r
r += 1
put_label(ws, r, 1, "Operator fees (per operating store)", bold=True)
for i in range(11):
    col = 2 + i
    f = f"=-{A['op_fee']}*{get_column_letter(col)}{stores_row}"
    put_calc(ws, r, col, f, fmt=FMT_D)
fee_r = r
r += 2

put_label(ws, r, 1, "Operating result", bold=True)
for i in range(11):
    col = 2 + i
    f = (f"={get_column_letter(col)}{net_rev_row}+{get_column_letter(col)}{cogs_r}"
         f"+{get_column_letter(col)}{labor_r}+{get_column_letter(col)}{oth_r}"
         f"+{get_column_letter(col)}{fee_r}")
    put_total(ws, r, col, f, fmt=FMT_D)
opres_r = r
r += 1
put_label(ws, r, 1, "Net cash flow (incl. capex)", bold=True)
for i in range(11):
    col = 2 + i
    f = f"={get_column_letter(col)}{opres_r}+{get_column_letter(col)}{capex_row}"
    put_total(ws, r, col, f, fmt=FMT_D)
ncf_r = r
r += 2

put_label(ws, r, 1, "10-yr total net cash flow", bold=True)
c = ws.cell(row=r, column=2,
            value=f"=SUM({get_column_letter(2)}{ncf_r}:{get_column_letter(12)}{ncf_r})")
c.font = total_font; c.fill = total_fill; c.number_format = FMT_D
r += 1
put_label(ws, r, 1, "10-yr NPV @ municipal discount rate", bold=True)
c = ws.cell(row=r, column=2,
            value=f"=NPV({A['disc_rate']},{get_column_letter(3)}{ncf_r}:{get_column_letter(12)}{ncf_r})+{get_column_letter(2)}{ncf_r}")
c.font = total_font; c.fill = total_fill; c.number_format = FMT_D
r += 1
put_label(ws, r, 1, "Total capex (memo)", bold=True)
c = ws.cell(row=r, column=2,
            value=f"=-SUM({get_column_letter(2)}{capex_row}:{get_column_letter(12)}{capex_row})")
c.font = total_font; c.fill = total_fill; c.number_format = FMT_D
r += 1
put_label(ws, r, 1, "Total operating losses over 10yrs", bold=True)
c = ws.cell(row=r, column=2,
            value=f"=MIN(0,SUM({get_column_letter(2)}{opres_r}:{get_column_letter(12)}{opres_r}))")
c.font = total_font; c.fill = total_fill; c.number_format = FMT_D


# =============================================================================
# 5. SENSITIVITY
# =============================================================================
ws = wb.create_sheet("Sensitivity")
ws.column_dimensions["A"].width = 32
for col in range(2, 14):
    ws.column_dimensions[get_column_letter(col)].width = 12

put_header(ws, 1, "Sensitivity — single-store steady-state operating result", span=12)
put_label(ws, 2, 1,
          "Rows: sales per sqft. Columns: discount to shoppers. Cell = steady-state annual operating result per store.",
          note=True)
put_label(ws, 3, 1,
          "Positive = profit, negative = annual subsidy needed. Read the crossover to see where the plan pencils.",
          note=True)

sales_per_sqft = [400, 450, 500, 550, 600, 700, 800, 900, 1019]
discounts = [0.05, 0.075, 0.10, 0.125, 0.15, 0.20, 0.25, 0.30]

r = 5
put_label(ws, r, 1, "Sales / sqft  ↓    Discount →", bold=True)
for j, d in enumerate(discounts):
    c = ws.cell(row=r, column=2 + j, value=d)
    c.font = label_bold
    c.number_format = FMT_PCT
    c.alignment = Alignment(horizontal="right")
r += 1

# steady-state per-store operating result =
# gross_revenue - shopper_discount - costs - op_fee
# = sqft * rev_sqft * [(1 - discount) - (cogs + labor + opex)] - op_fee
# Costs are per-unit (driven by gross rev); discount is a price cut that reduces net rev only.
grid_first_row = r
for i, s in enumerate(sales_per_sqft):
    put_label(ws, r, 1, f"${s:,}/sqft", bold=True)
    for j, d in enumerate(discounts):
        col = 2 + j
        f = (f"={A['sqft']}*{s}*((1-{d})-({A['cogs_pct']}+{A['labor_pct']}+{A['opex_pct']}))"
             f"-{A['op_fee']}")
        cell = ws.cell(row=r, column=col, value=f)
        cell.number_format = FMT_D
        cell.font = label_font
    r += 1

# Highlight negative results only — conditional, so the fill tracks live edits
grid_range = (f"B{grid_first_row}:"
              f"{get_column_letter(1 + len(discounts))}{grid_first_row + len(sales_per_sqft) - 1}")
ws.conditional_formatting.add(
    grid_range,
    CellIsRule(operator="lessThan", formula=["0"],
               fill=PatternFill("solid", start_color="FDEDEC", end_color="FDEDEC")),
)

r += 1
put_label(ws, r, 1, "Reading the table", bold=True, note=False); r += 1
notes = [
    "Industry avg net margin is ~1.6% (FMI 2024). At this model's April 10% discount assumption, every realistic sales/sqft produces a loss — because the discount exceeds the margin six times over.",
    "The store only clears profit with a small discount (≤5%) AND sales/sqft far above industry ($2,200+) — a combination unlikely for a new small-format municipal store.",
    "At the plausible $400–600/sqft for a new small-format store + 10% discount, each store needs $508k–$562k/yr in operating subsidy (the grid's own 10% column).",
    "Eliminating the operator fee improves the result by $400k/yr. The plan structure, however, requires a private operator.",
    "The 20–30% columns were added August 2026 after the administration announced a 30% core-basket discount (July 27, 2026). "
    "Above a ~7% discount the contribution margin is negative, so higher sales/sqft deepens the loss — the gradient flips; "
    "by a ~17% blended discount it is worse than −10¢ per dollar of sales. See the July 2026 Update tab for the blended scenarios.",
]
for n in notes:
    c = ws.cell(row=r, column=1, value=f"• {n}")
    c.font = label_font
    c.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    r += 1


# =============================================================================
# 5b. JULY 2026 UPDATE — the announced 30% core-basket discount
# =============================================================================
ws = wb.create_sheet("July 2026 Update")
ws.column_dimensions["A"].width = 48
for col in range(2, 8):
    ws.column_dimensions[get_column_letter(col)].width = 14
ws.column_dimensions["H"].width = 70

put_header(ws, 1, "July 2026 update — the announced 30% core-basket discount", span=8)
put_label(ws, 2, 1,
          "Added August 2026. The original tabs keep the April 2026 central case (10% discount). The April "
          "piece's tables now match this file only where that piece's banner says so — an August 2026 re-audit "
          "corrected the school-meals, FRESH, and SNAP-count inputs and rebuilt the Alternatives tab. "
          "Everything on this tab is additive; yellow cells are editable.",
          note=True)

U = {}  # name -> absolute address on this tab


def add_u(key, row, col=2):
    U[key] = f"'July 2026 Update'!${get_column_letter(col)}${row}"


r = 4
put_section(ws, r, "What was announced (July 27, 2026)", span=8); r += 1
put_label(ws, r, 1, "Core-basket discount vs. typical retail", bold=True)
put_input(ws, r, 2, 0.30, fmt=FMT_PCT0,
          comment="Mayor's Office, July 27, 2026. Vision plan wording is the operative "
                  "one: basket items 'priced ON AVERAGE 30 percent below market prices' "
                  "— a basket average, item-level variance allowed; no reference price "
                  "index disclosed. Basket = all fresh produce, meats, seafood + select "
                  "dairy, shelf-stable and frozen items (release: 'roughly 20 additional "
                  "categories'). Pricing cadence is the RFP's, not the vision plan's: "
                  "the RFP (pp. 12–13, 28) makes prices 'subject to monthly adjustments' "
                  "and uniform across all five stores; the vision plan itself says "
                  "prices change 'only periodically.' Non-basket items must be 'fairly "
                  "priced' too — the operator can't recover the discount with margin "
                  "elsewhere in the store. No income test. See "
                  "IMPLEMENTATION_LEDGER_2026-08.md.")
put_label(ws, r, 8, "Basket: all produce, meat, seafood + ~20 staple/dairy/refrigerated categories. Monthly price adjustments per the RFP, uniform citywide; applies to all shoppers.", note=True)
add_u("basket_disc", r); r += 1
put_label(ws, r, 1, "City Hall projection: cut to average total grocery bill", bold=True)
put_input(ws, r, 2, 0.15, fmt=FMT_PCT0,
          comment="Quoted figure, not a computed one (memo only — nothing downstream "
                  "uses it). July 27, 2026 release: 'cut New Yorkers' average grocery "
                  "bill by 15 percent, about $90 a month, or roughly $1,000 a year.'")
put_label(ws, r, 8, "Release: 'cut New Yorkers' average grocery bill by 15 percent, about $90 a month, or roughly $1,000 a year.'", note=True); r += 1
put_label(ws, r, 1, "Promised savings per shopper ($/month)", bold=True)
put_input(ws, r, 2, 90, fmt=FMT_D,
          comment="City Hall's own projection, July 27, 2026 release.")
add_u("savings_mo", r); r += 1
put_label(ws, r, 1, "Promised savings per shopper ($/yr)", bold=True)
put_calc(ws, r, 2, f"=12*{U['savings_mo']}", fmt=FMT_D, bold=True)
add_u("savings_yr", r); r += 1
put_label(ws, r, 1, "Average selling sqft per store (RFP-consistent)", bold=True)
put_input(ws, r, 2, 13800, fmt=FMT_SQFT,
          comment="(15,000 Peninsula/Bronx unit per RFP Appendix D — explicitly "
                  "including elevator-accessible mezzanine storage — + 9,000 La "
                  "Marqueta + 3 × 15,000 per RFP Appendix E, which instructs BK/QNS/SI "
                  "bids to assume 'approximately 15,000 square feet of selling feet' [sic]) "
                  "÷ 5 = 13,800. Recentered August 2026 on the operative RFP; "
                  "supersedes the 11,800 built from the releases (20k Peninsula "
                  "marketing figure + 3 × 10k portal minimums). DEFINITIONAL CAVEAT: "
                  "the Bronx 15,000 includes mezzanine storage and the five figures "
                  "mix selling vs gross definitions, so the average is not pure "
                  "selling area. Variant rows under Grid B carry the 11,800 (20k "
                  "marketing reading) and 14,800 (20k + 9k + 3×15k) cases.")
put_label(ws, r, 8, "RFP App. D: Bronx unit ~15,000 SF incl. mezzanine storage. App. E: BK/QNS/SI bids assume ~15,000 selling sqft. La Marqueta 9,000 (by 2029); Peninsula opens end-2027.", note=True)
add_u("avg_sqft", r); r += 1
put_label(ws, r, 1, "Industry-average sales per sqft (FMI 2025)", bold=True)
put_input(ws, r, 2, 1019, fmt=FMT_D,
          comment="FMI 2025 Food Industry Facts (released July 7, 2026): $19.59 per "
                  "sqft per week × 52 ≈ $1,019. This is the demand-consistent central "
                  "traffic for the ANNOUNCED plan: a store priced 30% below market is "
                  "capacity-bound (the depth of the discount is the point, and it draws "
                  "the crowd), so industry-average traffic "
                  "is the floor of the plausible range, not an optimistic case. "
                  "High-volume urban formats run far higher (Trader Joe's ~$2,000/sqft). "
                  "The April model's $500 stays as the quiet-store scenario: a 10%-off "
                  "store draws no queue.")
add_u("rev_ind", r); r += 1
put_label(ws, r, 1, "Persons per household (NYC)", bold=True)
put_input(ws, r, 2, 2.48, fmt="0.00",
          comment="Census QuickFacts, 2020–2024 ACS: NYC average household size 2.48. "
                  "Applied uniformly to convert households to New Yorkers. NYC SNAP "
                  "households average ~1.66 persons (OTDA 2025), so this factor "
                  "overstates people reached for SNAP-like populations; families with "
                  "children run ~3+.")
add_u("pph", r); r += 1
put_label(ws, r, 1, "Capital budget (five stores)", bold=True)
put_calc(ws, r, 2, f"={A['total_capex']}", fmt=FMT_D, bold=True)
put_label(ws, r, 8, "Announced $70M. No operating-subsidy figure has been disclosed in any release.", note=True); r += 2

put_section(ws, r, "Blended in-store discount", span=8); r += 1
put_label(ws, r, 1, "Core-basket share of store sales", bold=True)
put_input(ws, r, 2, 0.60, fmt=FMT_PCT0,
          comment="No official sales-share figure exists; the record brackets it three "
                  "ways (ledger): EDC board says discounted items ≈40% of store "
                  "OFFERINGS (a SKU share, via NY1); an industry consultant estimates "
                  "70–80% of sales VOLUME (FoodNavigator); and City Hall's own savings "
                  "math (30% basket cut → 15% average-bill cut) implies ~50% of a "
                  "shopper's total SPEND. 60% of store sales stays the central case; "
                  "the scenario columns bracket 50–100% share (the 25% blend column "
                  "implies ~83%).")
add_u("basket_share", r); r += 1
put_label(ws, r, 1, "Central blended in-store discount", bold=True)
put_calc(ws, r, 2, f"={U['basket_disc']}*{U['basket_share']}", fmt=FMT_PCT, bold=True)
put_label(ws, r, 8, "30% × basket share. City Hall's own 15% average-bill figure is the conservative floor; 30% is the bound if everything in the store is basket.", note=True)
add_u("central_blend", r); r += 1
put_label(ws, r, 1, "Annuity factor — value today of $1/yr for 10 years", bold=True)
put_calc(ws, r, 2, f"=(1-(1+{A['disc_rate']})^(-10))/{A['disc_rate']}", fmt="0.000000", bold=True)
put_label(ws, r, 8, "At the Assumptions tab's 4% municipal rate. Used to state leverage 'in today's dollars': future discount dollars and future subsidy checks are both worth less than dollars today, while the $70M of construction is paid up front.", note=True)
add_u("af10", r); r += 2

# ---- Scenario grid A — quiet store ----
put_section(ws, r, "Scenario grid A — quiet store (the April model's $500/sqft traffic)", span=8); r += 1
scenario_headers = ["April (10%)", "15%", "Central blend", "20%", "25%", "30%"]
put_label(ws, r, 1, "Scenario", bold=True)
for j, h in enumerate(scenario_headers):
    c = ws.cell(row=r, column=2 + j, value=h)
    c.font = label_bold
    c.alignment = Alignment(horizontal="right")
r += 1
put_label(ws, r, 1, "Blended in-store discount", bold=True)
scenario_disc = [f"={A['discount']}", 0.15, f"={U['central_blend']}", 0.20, 0.25, 0.30]
for j, d in enumerate(scenario_disc):
    put_calc(ws, r, 2 + j, d, fmt=FMT_PCT)
u_d_row = r
put_label(ws, r, 8, "April column reads the Assumptions discount cell (10% central case), kept for apples-to-apples comparison.", note=True); r += 1
put_label(ws, r, 1, "Gross revenue per store ($/yr)", bold=True)
put_calc(ws, r, 2, f"={U['avg_sqft']}*{A['rev_sqft']}", fmt=FMT_D, bold=True)
put_label(ws, r, 8, "Same across scenarios: announced avg footprint × the Assumptions tab's $/sqft central case.", note=True)
u_gross_row = r; r += 1
put_label(ws, r, 1, "Contribution per $1 of gross sales", bold=True)
for j in range(6):
    col = 2 + j
    f = (f"=(1-{get_column_letter(col)}{u_d_row})"
         f"-({A['cogs_pct']}+{A['labor_pct']}+{A['opex_pct']})")
    put_calc(ws, r, col, f, fmt=FMT_PCT)
u_contrib_row = r
put_label(ws, r, 8, "At 30% off, basket items price at 70% of retail vs ~72% industry acquisition cost — at or below cost before any labor.", note=True); r += 1
put_label(ws, r, 1, "Operating result per store ($/yr)", bold=True)
for j in range(6):
    col = 2 + j
    f = f"=$B${u_gross_row}*{get_column_letter(col)}{u_contrib_row}-{A['op_fee']}"
    put_total(ws, r, col, f, fmt=FMT_D)
u_opres_row = r; r += 1
put_label(ws, r, 1, "Annual operating subsidy, five stores ($/yr)", bold=True)
for j in range(6):
    col = 2 + j
    put_calc(ws, r, col, f"=-5*{get_column_letter(col)}{u_opres_row}", fmt=FMT_D)
u_sub_row = r; r += 1
put_label(ws, r, 1, "Ten-year public envelope (capex + 10 × subsidy)", bold=True)
for j in range(6):
    col = 2 + j
    put_calc(ws, r, col, f"={A['total_capex']}+10*{get_column_letter(col)}{u_sub_row}", fmt=FMT_D)
u_env_row = r; r += 1
put_label(ws, r, 1, "Discount $ reaching shoppers, five stores ($/yr)", bold=True)
for j in range(6):
    col = 2 + j
    put_calc(ws, r, col, f"=5*$B${u_gross_row}*{get_column_letter(col)}{u_d_row}", fmt=FMT_D)
u_transfer_row = r; r += 1
put_label(ws, r, 1, "Households receiving the full promised savings", bold=True)
for j in range(6):
    col = 2 + j
    put_calc(ws, r, col, f"={get_column_letter(col)}{u_transfer_row}/{U['savings_yr']}", fmt=FMT_N)
u_hh_row = r; r += 1
put_label(ws, r, 1, "Nominal leverage = 10-yr benefit ÷ envelope (secondary)", bold=True)
for j in range(6):
    col = 2 + j
    put_calc(ws, r, col, f"=10*{get_column_letter(col)}{u_transfer_row}/{get_column_letter(col)}{u_env_row}", fmt='0.00"x"')
u_lev_row = r
put_label(ws, r, 8, "No time value — kept to reconcile with the sum-of-checks envelope. The PV row below is the headline metric.", note=True); r += 1
put_label(ws, r, 1, "Food benefit per public dollar, in today's dollars", bold=True)
for j in range(6):
    col = 2 + j
    f = (f"=({get_column_letter(col)}{u_transfer_row}*{U['af10']})"
         f"/({A['total_capex']}+{get_column_letter(col)}{u_sub_row}*{U['af10']})")
    put_total(ws, r, col, f, fmt='0.00"x"')
u_pvlev_row = r
put_label(ws, r, 8, "PV of ten transfer years ÷ (capex + PV of ten subsidy years), at the 4% muni rate. A 10%-off store draws no queue, so April's quiet-small-format traffic was the right frame then. Steady state, all five open; ramp and phasing ignored on both cost and benefit.", note=True); r += 2

# ---- Scenario grid B — demand-consistent ----
put_section(ws, r, "Scenario grid B — capacity-bound store at industry-average traffic (the announced plan's frame)", span=8); r += 1
put_label(ws, r, 1, "Gross revenue per store ($/yr)", bold=True)
put_calc(ws, r, 2, f"={U['avg_sqft']}*{U['rev_ind']}", fmt=FMT_D, bold=True)
put_label(ws, r, 8, "A store priced 30% below market is capacity-bound — the depth of the discount is the point, and it draws the crowd. Industry-average traffic is the FLOOR of the busy range (Trader Joe's runs ~$2,000/sqft), so this frame is conservative on cost and generous on reach at once.", note=True)
c_gross_row = r; r += 1
put_label(ws, r, 1, "Operating result per store ($/yr)", bold=True)
for j in range(6):
    col = 2 + j
    f = f"=$B${c_gross_row}*{get_column_letter(col)}{u_contrib_row}-{A['op_fee']}"
    put_total(ws, r, col, f, fmt=FMT_D)
c_opres_row = r
put_label(ws, r, 8, "Once contribution is negative, busier is worse — the discount makes success expensive.", note=True); r += 1
put_label(ws, r, 1, "Annual operating subsidy, five stores ($/yr)", bold=True)
for j in range(6):
    col = 2 + j
    put_calc(ws, r, col, f"=-5*{get_column_letter(col)}{c_opres_row}", fmt=FMT_D)
c_sub_row = r; r += 1
put_label(ws, r, 1, "Ten-year public envelope (capex + 10 × subsidy)", bold=True)
for j in range(6):
    col = 2 + j
    put_total(ws, r, col, f"={A['total_capex']}+10*{get_column_letter(col)}{c_sub_row}", fmt=FMT_D)
c_env_row = r; r += 1
put_label(ws, r, 1, "Discount $ reaching shoppers, five stores ($/yr)", bold=True)
for j in range(6):
    col = 2 + j
    put_calc(ws, r, col, f"=5*$B${c_gross_row}*{get_column_letter(col)}{u_d_row}", fmt=FMT_D)
c_transfer_row = r; r += 1
put_label(ws, r, 1, "Households receiving the full promised savings", bold=True)
for j in range(6):
    col = 2 + j
    put_total(ws, r, col, f"={get_column_letter(col)}{c_transfer_row}/{U['savings_yr']}", fmt=FMT_N)
c_hh_row = r
put_label(ws, r, 8, "Transfer capacity ÷ the promised $90/month (= $1,080/yr; the release rounds to 'roughly $1,000 a year'). The stores can be generous or broad, not both.", note=True); r += 1
put_label(ws, r, 1, "New Yorkers (households × persons per HH)", bold=True)
for j in range(6):
    col = 2 + j
    put_calc(ws, r, col, f"={get_column_letter(col)}{c_hh_row}*{U['pph']}", fmt=FMT_N)
c_ppl_row = r; r += 1
put_label(ws, r, 1, "Nominal leverage = 10-yr benefit ÷ envelope (secondary)", bold=True)
for j in range(6):
    col = 2 + j
    put_calc(ws, r, col, f"=10*{get_column_letter(col)}{c_transfer_row}/{get_column_letter(col)}{c_env_row}", fmt='0.00"x"')
c_lev_row = r
put_label(ws, r, 8, "No time value — kept to reconcile with the sum-of-checks envelope. The PV row below is the headline metric.", note=True); r += 1
put_label(ws, r, 1, "Food benefit per public dollar, in today's dollars", bold=True)
for j in range(6):
    col = 2 + j
    f = (f"=({get_column_letter(col)}{c_transfer_row}*{U['af10']})"
         f"/({A['total_capex']}+{get_column_letter(col)}{c_sub_row}*{U['af10']})")
    put_total(ws, r, col, f, fmt='0.00"x"')
c_pvlev_row = r
put_label(ws, r, 8, "PV of ten transfer years ÷ (capex + PV of ten subsidy years), at the 4% muni rate. Central: 0.69x.", note=True); r += 1
put_label(ws, r, 1, "— envelope if the footprint reads 11,800 sqft (20k marketing reading)", bold=True)
put_calc(ws, r, 2,
         f"={A['total_capex']}+50*((11800*{U['rev_ind']})*(D{u_d_row}-(1-({A['cogs_pct']}+{A['labor_pct']}+{A['opex_pct']})))+{A['op_fee']})",
         fmt=FMT_D)
put_label(ws, r, 8, "The May 18 release and the EDC program page carry 20,000 sqft for the Peninsula where the RFP says ~15,000 incl. mezzanine; 20k + 9k La Marqueta + 3 × 10k portal minimums gives avg 11,800 — the marketing reading. Envelope ~$156M.", note=True); r += 1
put_label(ws, r, 1, "— envelope if the footprint reads 14,800 sqft (20k Peninsula + App. E 15k)", bold=True)
put_calc(ws, r, 2,
         f"={A['total_capex']}+50*((14800*{U['rev_ind']})*(D{u_d_row}-(1-({A['cogs_pct']}+{A['labor_pct']}+{A['opex_pct']})))+{A['op_fee']})",
         fmt=FMT_D)
put_label(ws, r, 8, "(20,000 + 9,000 + 3 × 15,000) ÷ 5 = 14,800 — keeping the release's Peninsula figure alongside the RFP's Appendix E instruction. Envelope ~$173M. The $156M–$173M span brackets every official footprint reading around the $167M central.", note=True); r += 1
put_label(ws, r, 1, "Bounding case: city pays the full discount", bold=True)
put_calc(ws, r, 2, f"={A['total_capex']}+10*D{c_transfer_row}", fmt=FMT_D)
put_label(ws, r, 8, "If the city reimbursed the full discount rather than the operating deficit, the central envelope would run ~$197M. This is NOT the RFP's structure: the RFP is deficit-based (p. 7 — affordability payments cover 'operating deficits driven by the sale of Core Basket SKUs at discounted prices,' net of the covered rent and property taxes), which corroborates this model's loss-based central. Fee convention: the operator fee is treated as absorbed within the discount reimbursement, not added on top.", note=True); r += 2

# ---- The two central cases ----
put_section(ws, r, "The two central cases — the demand-consistent diagonal", span=8); r += 1
put_label(ws, r, 1, "", bold=True)
for j, h in enumerate(["April model\n(10% × $500)", "Announced plan\n(18% blend × $1,019)",
                       "Change (vs published\n$100M base)"]):
    c = ws.cell(row=r, column=2 + j, value=h)
    c.font = label_bold
    c.alignment = Alignment(horizontal="right", wrap_text=True)
ws.row_dimensions[r].height = 26
r += 1
diag_rows = [
    ("Ten-year public envelope (nominal — the sum of checks)", u_env_row, c_env_row, FMT_D),
    ("Annual operating subsidy, five stores", u_sub_row, c_sub_row, FMT_D),
    ("Discount $ reaching shoppers ($/yr)", u_transfer_row, c_transfer_row, FMT_D),
    ("Households getting the full deal", u_hh_row, c_hh_row, FMT_N),
    ("Food benefit per public dollar, in today's dollars", u_pvlev_row, c_pvlev_row, '0.00"x"'),
    ("Nominal leverage (secondary)", u_lev_row, c_lev_row, '0.00"x"'),
]
for label, april_row, ann_row, fmt in diag_rows:
    put_label(ws, r, 1, label, bold=True)
    put_calc(ws, r, 2, f"=B{april_row}", fmt=fmt)
    put_calc(ws, r, 3, f"=D{ann_row}", fmt=fmt, bold=True)
    put_calc(ws, r, 4, f"=D{ann_row}/B{april_row}-1", fmt="+0%;-0%")
    if label.startswith("Ten-year public envelope"):
        diag_env_row = r
    elif label.startswith("Annual operating subsidy"):
        diag_sub_row = r
    elif label.startswith("Discount $ reaching shoppers"):
        diag_transfer_row = r
    r += 1
put_label(ws, r, 1, "New Yorkers getting the full deal (announced)", bold=True)
put_calc(ws, r, 3, f"=D{c_ppl_row}", fmt=FMT_N, bold=True)
put_label(ws, r, 8, "The piece's headline numbers come from this block: the bill is the nominal envelope (the checks the city writes); the per-dollar metric is the PV row, branded 'in today's dollars.' April column = grid A at the Assumptions 10%; announced = grid B at the central blend. The April column at this footprint lands at $100.35M ≈ the published $100M, so growth statements use the published base: the bill grows about two-thirds; benefit per dollar in today's dollars more than doubles (×2.33).", note=True); r += 2

# ---- Same money, other shapes ----
put_section(ws, r, "Same money, other shapes — the announced envelope spent differently", span=8); r += 1
put_label(ws, r, 1, "Annual pool (announced envelope ÷ 10 yrs)", bold=True)
put_calc(ws, r, 2, f"=C{diag_env_row}/10", fmt=FMT_D, bold=True)
pool_row = r; r += 1
put_label(ws, r, 1, "Warehouse-club membership utilization", bold=True)
put_calc(ws, r, 2, f"={A['warehouse_util']}", fmt=FMT_PCT0, bold=True)
put_label(ws, r, 8, "Reads the Assumptions tab's utilization input (author assumption, 70%) — the single source of truth shared with the Alternatives tab's Costco/BJ's row.", note=True)
add_u("wh_util", r); r += 1
shape_rows = [
    ("Cash at the stores' own depth ($1,080/HH/yr)",
     f"=B{pool_row}/{U['savings_yr']}",
     "The depth-matched comparison: same full deal, no buildings. The stores lose even at their own depth."),
    ("SNAP-style supplement ($20/mo = $240/HH/yr)",
     f"=B{pool_row}/({A['snap_topup']}*12)",
     "Direct transfer at the Alternatives tab's supplement depth."),
    ("Warehouse-club memberships funded ($65 each)",
     f"=B{pool_row}/{A['warehouse_fee']}",
     "Memberships bought outright each year."),
    ("— actively used (× utilization)",
     f"=B{pool_row}/{A['warehouse_fee']}*{U['wh_util']}",
     "Households actually saving (~$600/yr realized, blended car/transit access — Alternatives tab)."),
]
shape_first_row = r
for label, formula, note in shape_rows:
    put_label(ws, r, 1, label, bold=True)
    put_calc(ws, r, 2, formula, fmt=FMT_N, bold=True)
    put_calc(ws, r, 3, f"=B{r}*{U['pph']}", fmt=FMT_N)
    put_label(ws, r, 8, note, note=True)
    r += 1
put_label(ws, r, 1, "(column C = New Yorkers, × persons per household)", note=True); r += 1
put_label(ws, r, 1, "FRESH-style abatements: new supermarkets (envelope ÷ $1.08M public cost per store — Comptroller Fiscal Note 4-2024)", bold=True)
put_calc(ws, r, 2, f"=C{diag_env_row}/{A['fresh_per_store']}", fmt=FMT_1, bold=True)
put_label(ws, r, 8, "≈155 stores over the ten years — about 15/yr vs the program's historic ~2/yr; treat as a ceiling on uptake, not a forecast.", note=True)
fresh_stores_row = r; r += 1
put_label(ws, r, 1, "— households beyond a half-mile walk to a supermarket (USDA, all incomes)", bold=True)
put_input(ws, r, 2, 143000, fmt=FMT_N,
          comment="USDA Food Access Research Atlas, half-mile distance test, all "
                  "incomes — adopted as the right instrument for a walking city. "
                  "The one-mile standard, calibrated for driving suburbs, returns "
                  "~0 for NYC and is rejected. The low-income core of this "
                  "population is roughly 32,000 households (~80,000 people). "
                  "≈355,000 New Yorkers at the model's 2.48 persons per household "
                  "(FARA's own person count for the same tracts is ~377,000).")
put_calc(ws, r, 3, f"=B{r}*{U['pph']}", fmt=FMT_N)
put_label(ws, r, 8, "Access benefit, not a transfer: the full ~$300/HH/yr is reserved for the low-income core (~32,000 HH), tapering across the rest — see the update methodology. The raw walking-distance catchment (~16,100 HH/store at 2.48 persons/HH, overlapping) is a ceiling on siting, not need.", note=True)
add_u("low_access_hh", r); r += 1
put_label(ws, r, 1, "Bodega-upgrade route: reach ceiling (deduped)", bold=True)
put_input(ws, r, 2, 3334088, fmt=FMT_N,
          comment="Census QuickFacts 2020–2024: 3,334,088 NYC households. The April "
                  "table's 5M HH (10,000 bodegas × 500 HH walking distance) "
                  "double-counts overlapping catchments — nearly every NYC "
                  "household lives near a bodega, so citywide households is the "
                  "honest ceiling.")
put_label(ws, r, 8, "The $167M envelope is ~2.2× the bodega program's $76M natural scope (10k bodegas × $7.6k). The extra money buys depth and duration (produce support beyond 2 years), not reach, so leverage runs ~10–20× depending on how benefit scales — below the April table's original 32.9×, corrected to ~10–20× in both pieces, August 2026.", note=True)
add_u("bodega_ceiling", r); r += 1
put_label(ws, r, 1, "Memo: SNAP households in NYC (Assumptions)", bold=True)
put_calc(ws, r, 2, f"={A['snap_households']}", fmt=FMT_N)
put_label(ws, r, 8, "The stores' full deal reaches ~1.1% of NYC SNAP households.", note=True); r += 1
put_label(ws, r, 1, "Implied basket spend at the store for the full deal ($/yr)", bold=True)
put_calc(ws, r, 2, f"={U['savings_yr']}/{U['basket_disc']}", fmt=FMT_D, bold=True)
put_label(ws, r, 8, "≈ $300/month of core-basket purchases at the municipal store — effectively a household's whole fresh-and-staples shop. Consistent with City Hall's arithmetic: $90 ≈ 15% of a $600/month total bill.", note=True); r += 2

# ---- Horizon — does a longer window rescue it? ----
put_section(ws, r, "Horizon — does a longer window rescue it?", span=8); r += 1
put_label(ws, r, 1, "Operating leverage (transfer ÷ subsidy, steady state)", bold=True)
put_calc(ws, r, 2, f"=C{diag_transfer_row}/C{diag_sub_row}", fmt='0.00"x"', bold=True)
put_label(ws, r, 8, "Excluding construction, the running operation converts subsidy to shopper benefit at 1.3:1.", note=True); r += 1
put_label(ws, r, 1, "Envelope in today's dollars (construction + PV of ten subsidy years)", bold=True)
put_calc(ws, r, 2, f"={A['total_capex']}+C{diag_sub_row}*{U['af10']}", fmt=FMT_D, bold=True)
put_label(ws, r, 8, "≈$149M at the 4% muni rate — the same checks the nominal $167M envelope counts, discounted to today.", note=True); r += 1
put_label(ws, r, 1, "Phased-calendar variant (stores open 1/2/2 across years 1–3; 44 store-years)", bold=True)
put_calc(ws, r, 2, f"={A['total_capex']}+44*(-D{c_opres_row})", fmt=FMT_D, bold=True)
put_label(ws, r, 8, "Nominal envelope if openings phase over 2027–29: 44 store-years of the central steady-state loss instead of 50; ramp ignored.", note=True); r += 1
put_label(ws, r, 1, "30-year leverage, today's dollars", bold=True)
put_calc(ws, r, 2,
         f"=(C{diag_transfer_row}*((1-(1+{A['disc_rate']})^(-30))/{A['disc_rate']}))"
         f"/({A['total_capex']}+C{diag_sub_row}*((1-(1+{A['disc_rate']})^(-30))/{A['disc_rate']}))",
         fmt='0.00"x"', bold=True)
put_label(ws, r, 8, "Tripling the window still doesn't clear the floor: thirty subsidized years, discounted at the Assumptions tab's rate, deliver less than a dollar of benefit per public dollar.", note=True); r += 1
put_label(ws, r, 1, "Run forever (perpetuity at the 4% rate)", bold=True)
put_calc(ws, r, 2, f"=(C{diag_transfer_row}/{A['disc_rate']})/({A['total_capex']}+C{diag_sub_row}/{A['disc_rate']})", fmt='0.00"x"', bold=True)
put_label(ws, r, 8, "The NAIVE forever figure: flat nominal flows, original fit-out never replaced. Kept for reference — it is unstable in both directions (see the Asset value and Real rates blocks below): indexing the flows to food prices pushes it up to ~1.16, and charging fit-out replacement (user cost) pulls it to 0.70–0.82, under the cash floor at every rate in the band. The old ~4.2% break-even framing is retired (economist-review round, Aug 2026).", note=True); r += 2

# ---- Asset value & symmetry — the economist-review round (added Aug 2026) ----
put_section(ws, r, "Asset value — suppose the city sells the stores (added Aug 2026)", span=8); r += 1
env_pv_f = f"({A['total_capex']}+C{diag_sub_row}*{U['af10']})"
ben_pv_f = f"(C{diag_transfer_row}*{U['af10']})"
put_label(ws, r, 1, "Year-ten sale value — researched low ($)", bold=True)
put_input(ws, r, 2, 3_000_000, fmt=FMT_D,
          comment="Researched components (research/economist-review_2026-08/RESEARCH_NOTES.md, "
                  "2026-08-06): equipment ≈$0 at year ten (DOE 2024 CRE final rule: 10-yr "
                  "life in buildings >5,000 sqft, 14-yr average across sizes; used "
                  "food-retail equipment liquidates at 10–30 cents on the dollar; 5-yr "
                  "MACRS, Rev. Proc. 87-56 class 57.0). Fit-out in leased space reverts "
                  "to the landlord at the RFP's 10-yr base term. La Marqueta building at "
                  "sold-retail market value: STNL sold median $309/sqft (Colliers 1H2025) "
                  "× 9,000 sqft ≈ $2.8M. Low case = the building at the sold median only.")
add_u("tv_low", r); r += 1
put_label(ws, r, 1, "Year-ten sale value — researched high ($)", bold=True)
put_input(ws, r, 2, 12_000_000, fmt=FMT_D,
          comment="High case = La Marqueta near its NYC replacement-cost cap "
                  "($600–1,200/sqft × 9,000 ≈ $5.4–10.8M) plus equipment salvage. The "
                  "$30M row below is the GENEROUS bound (the full appropriation recovered "
                  "at cost); the dark-store litigation record documents purpose-built "
                  "retail routinely selling 50%+ below construction cost.")
add_u("tv_high", r); r += 1
put_label(ws, r, 1, "Leverage with a year-ten sale credited — researched low", bold=True)
put_calc(ws, r, 2,
         f"={ben_pv_f}/({env_pv_f}-{U['tv_low']}*(1+{A['disc_rate']})^(-10))",
         fmt='0.00"x"')
r += 1
put_label(ws, r, 1, "Leverage with a year-ten sale credited — researched high", bold=True)
put_calc(ws, r, 2,
         f"={ben_pv_f}/({env_pv_f}-{U['tv_high']}*(1+{A['disc_rate']})^(-10))",
         fmt='0.00"x"')
put_label(ws, r, 8, "The researched band moves the verdict from 0.69 to 0.70–0.73.", note=True); r += 1
put_label(ws, r, 1, "Leverage if the sale recovers $30M (generous bound)", bold=True)
put_calc(ws, r, 2,
         f"={ben_pv_f}/({env_pv_f}-30000000*(1+{A['disc_rate']})^(-10))",
         fmt='0.00"x"')
put_label(ws, r, 8, "Even the whole La Marqueta appropriation back at cost reads 0.80.", note=True); r += 1
put_label(ws, r, 1, "Sale value needed to reach the 1.0x cash floor ($)", bold=True)
put_calc(ws, r, 2,
         f"=({env_pv_f}-{ben_pv_f})*(1+{A['disc_rate']})^10",
         fmt=FMT_D, bold=True)
put_label(ws, r, 8, "≈$68.5M — essentially the full $70M back. Crediting the FULL $70M reproduces the perpetuity row above EXACTLY. That is an identity, not a coincidence: with a sale at the whole capex K, the ratio collapses to (transfer ÷ rate) ÷ (capex + subsidy ÷ rate) at any horizon and any rate. 'Sell it for what you paid' and 'run it forever' are the same claim — an asset that never wears out.", note=True); r += 1
put_label(ws, r, 1, "Market rent, blended ($/sqft/yr, researched)", bold=True)
put_input(ws, r, 2, 40, fmt=FMT_D,
          comment="$25–60 band, blended ~$40 by RFP square footage: Peninsula-complex "
                  "retail asking $25/sqft (PropertyShark, Aug 2026); E 116th St comps "
                  "$21 (second floor) to $80 (small ground-floor) with a big-box anchor "
                  "discount; C&W NYC-metro shopping-center average $37.32 (Q4 2025). "
                  "Supersedes the April memo's unsourced $50 for THIS block; the April "
                  "Assumptions memo cell is untouched (memo-only there).")
add_u("rent_psf", r); r += 1
put_label(ws, r, 1, "Implied market rent, five stores ($/yr)", bold=True)
put_calc(ws, r, 2, f"=5*{U['avg_sqft']}*{U['rent_psf']}", fmt=FMT_D)
add_u("rent_5", r); r += 1
put_label(ws, r, 1, "Full-balance-sheet leverage — rent charged, low sale credited", bold=True)
put_calc(ws, r, 2,
         f"={ben_pv_f}/({env_pv_f}+{U['rent_5']}*{U['af10']}-{U['tv_low']}*(1+{A['disc_rate']})^(-10))",
         fmt='0.00"x"')
r += 1
put_label(ws, r, 1, "Full-balance-sheet leverage — rent charged, high sale credited", bold=True)
put_calc(ws, r, 2,
         f"={ben_pv_f}/({env_pv_f}+{U['rent_5']}*{U['af10']}-{U['tv_high']}*(1+{A['disc_rate']})^(-10))",
         fmt='0.00"x"')
put_label(ws, r, 8, "The model excludes rent and property taxes (city-covered) AND credits no sale — crediting the sale while never charging for the space would be the one incoherent accounting. Do both and the ratio reads 0.61–0.63, at or below the published 0.69: the balance-sheet correction does not rescue the program.", note=True); r += 1
put_label(ws, r, 1, "Leased-site contingency: three 15,000-sqft leases at researched rent ($/yr)", bold=True)
put_calc(ws, r, 2, f"=3*15000*{U['rent_psf']}", fmt=FMT_D)
add_u("leased_rent", r); r += 1
put_label(ws, r, 1, "Leverage if those rent checks join the bill", bold=True)
put_calc(ws, r, 2,
         f"={ben_pv_f}/({env_pv_f}+{U['leased_rent']}*{U['af10']})",
         fmt='0.00"x"')
put_label(ws, r, 8, "The RFP commits NYCEDC to 'pay all rent and property taxes (as applicable)' at all five sites (p. 3) and 'may lease space from a private landlord' for BK/QNS/SI (p. 9). If it does, the rent checks belong in the bill and are NOT in the $167M: $11–27M nominal over the decade at the $25–60 band, ratio ~0.60–0.65. Labeled contingency, not a projection — EDC says it prioritizes city-owned sites; the first Q&A round (due Aug 14, 2026) may resolve tenure.", note=True); r += 2

put_section(ws, r, "Real rates and replacement — the honest run-forever (added Aug 2026)", span=8); r += 1
put_label(ws, r, 1, "Inflation for indexed flows (10-yr breakeven)", bold=True)
put_input(ws, r, 2, 0.0226, fmt="0.00%",
          comment="FRED T10YIE, 2026-08-06: 2.26%. Food-at-home CPI +2.7% y/y through "
                  "June 2026 (USDA ERS). The transfer and the Affordability Payments are "
                  "shares of nominal food sales, so they index with food prices; "
                  "discounting indexed flows at the nominal rate is equivalent to "
                  "discounting flat flows at the implied real rate. Context: NYC GO "
                  "10-yr tax-exempt priced at 3.61% on Aug 6, 2026 (Fiscal 2027 Series A "
                  "official statement), so the model's 4% pin is if anything high "
                  "against the city's actual 10-yr borrowing cost.")
add_u("g_infl", r); r += 1
put_label(ws, r, 1, "Implied real rate ((1 + 4%) ÷ (1 + inflation) − 1)", bold=True)
put_calc(ws, r, 2, f"=(1+{A['disc_rate']})/(1+{U['g_infl']})-1", fmt="0.00%")
add_u("r_real", r); r += 1
put_label(ws, r, 1, "Annuity factor at the real rate (10 years)", bold=True)
put_calc(ws, r, 2, f"=(1-(1+{U['r_real']})^(-10))/{U['r_real']}", fmt="0.000000")
add_u("af_real", r); r += 1
put_label(ws, r, 1, "Ten-year leverage, indexed flows (real-rate read)", bold=True)
put_calc(ws, r, 2,
         f"=(C{diag_transfer_row}*{U['af_real']})/({A['total_capex']}+C{diag_sub_row}*{U['af_real']})",
         fmt='0.00"x"')
put_label(ws, r, 8, "0.72–0.73 across inflation 2–2.5%. Direction: favorable to the program vs the published 0.689 — disclosed, not re-based. The flat-nominal-at-4% convention understates the program's ratio, which is the conservative direction for this piece's verdict.", note=True); r += 1
put_label(ws, r, 1, "No-replacement perpetuity, indexed flows (naive)", bold=True)
put_calc(ws, r, 2,
         f"=C{diag_transfer_row}/({A['total_capex']}*({A['disc_rate']}-{U['g_infl']})+C{diag_sub_row})",
         fmt='0.00"x"')
put_label(ws, r, 8, "≈1.16 — the published 1.01 pushed UP by indexing. The naive forever figure is unstable in both directions; the user-cost row below is the honest one.", note=True); r += 1
put_label(ws, r, 1, "Fit-out economic life (years, author central)", bold=True)
put_input(ws, r, 2, 12.5, fmt="0.0",
          comment="Band 10–15 years: DOE 2024 CRE rule uses 10 years in >5,000 sqft "
                  "buildings (14-yr average across sizes); qualified improvement "
                  "property carries a 15-yr recovery period; trade sources put "
                  "commercial refrigeration at 10–15. Running the stores forever means "
                  "replacing this fit-out forever.")
add_u("fitout_life", r); r += 1
put_label(ws, r, 1, "User cost of capital, K × (real rate + 1 ÷ life) ($/yr)", bold=True)
put_calc(ws, r, 2, f"={A['total_capex']}*({U['r_real']}+1/{U['fitout_life']})", fmt=FMT_D)
add_u("user_cost", r); r += 1
put_label(ws, r, 1, "Run forever with replacement charged (user cost)", bold=True)
put_calc(ws, r, 2,
         f"=C{diag_transfer_row}/(C{diag_sub_row}+{U['user_cost']})",
         fmt='0.00"x"', bold=True)
put_label(ws, r, 8, "Band 0.70–0.82 across real rate 1.46–1.96% × life 10–15 years — UNDER the cash floor at every combination, no knife edge. This replaces the old ~4.2% break-even story, which was an artifact of flat nominal flows and never replacing the shelves. Hall–Jorgenson user cost; annualized capital charges are a recognized equivalent of NPV treatment (EPA Guidelines for Preparing Economic Analyses §6.1.2).", note=True); r += 2

update_notes = [
    "The demand-consistent diagonal: April's central was (10% discount, $500/sqft) — a modest discount at quiet "
    "small-format traffic. The announced plan's central is an 18% blend at $1,019/sqft (FMI 2025) on the "
    "RFP-consistent 13,800 sqft average footprint — a store priced 30% below market is capacity-bound, so "
    "industry-average traffic is the floor of the plausible range. This is generous to the plan on the leverage "
    "metric (busier stores score higher), which is what makes it fair.",
    "Quiet-store fallback: if the announced stores somehow stay at $500/sqft, the envelope falls to ~$128M and "
    "today's-dollars leverage to 0.43× — one branch makes the plan more expensive, the other less effective; "
    "neither approaches the 1.0x floor.",
    "No official operating-subsidy figure exists anywhere in the public record as of August 2026 — the vision plan "
    "contains zero dollar figures, and the adopted FY2027 budget, Schedule C, Comptroller testimony, and IBO are all "
    "silent. The subsidy MECHANISM is official (annual affordability payments, sized by operators' own subsidy bids, "
    "selection early-to-spring 2027), so this tab's envelope is an estimate of a number the city has not yet priced. "
    "Full ledger of the official record: IMPLEMENTATION_LEDGER_2026-08.md.",
    "Apples-to-apples: the same throughput-capped frame applied to the April 10% case and the announced scenarios. "
    "The Alternatives tab's 0.6x / 1.3x plan rows (matching the April piece's published table) use a more generous "
    "reach × spend frame; this frame caps benefit at what the stores can physically sell, which is the binding "
    "constraint at a 30% discount.",
    "Every additional point of discount is a 1:1 transfer at the margin (benefit and subsidy rise by the same dollars), "
    "so deepening the discount pulls leverage toward — never past — the 1.0x direct-transfer floor at any plausible "
    "volume. Along the VOLUME axis the story differs: crossing happens only above ~$2,200/sqft in today's dollars "
    "(≈2.2× the industry average) where the subsidy runs $19–37M/yr; as volume grows without bound, leverage "
    "approaches a ceiling of discount ÷ (discount − margin): 1.64× at the 18% blend, 1.30× at 30% — and the "
    "ceiling is cost-structure-dependent (95% costs → 1.38×; 91% → 2.0×). "
    "The $70M of construction and a decade of overhead hold the announced central to 0.69×.",
    "Cross-check: the only published outside per-store estimate of the announced plan (July 2026, a state comptroller "
    "candidate) applies 30% to ALL revenue on a ~$5.4M sales base (~$1.6M give-up, ~$1.5M loss per store). Two "
    "offsetting errors — a discount roughly double the blended reality, on a revenue base under half this model's "
    "central — land his bottom line in the same ballpark as this model's −$1.9M/store. Right ballpark, wrong "
    "mechanics. An earlier estimate (NY Post, April 17, 2026) put the East Harlem store's loss at at least "
    "$300,000/yr on pre-announcement economics.",
    "External anchor: DoD commissaries deliver ~24–25% patron savings (DeCA self-reported; GAO-17-80 criticized the "
    "methodology) on a $1.53–1.57B/yr appropriation (FY2024–26) — at face value roughly dollar-for-dollar leverage "
    "at 236-store scale, consistent with this model's ceiling logic.",
    "Sell-out risk (reasoning, not data): a 30% below-market price with no income test is a textbook shortage setup — "
    "demand rationed by queue, empty shelf, and per-customer limits; a 30% spread also makes resale arbitrage "
    "profitable. The RFP (July 27) requires operators to prevent excessive bulk purchases via a voluntary free "
    "savings card; numeric limits are not yet set.",
    "Asset value & symmetry (Aug 2026, after an economist reader's critique): crediting a year-ten sale at the "
    "researched $3–12M reads 0.70–0.73; the full $70M reproduces the perpetuity exactly (an identity). Charging "
    "market rent for the space while crediting the sale lands 0.61–0.63. The honest run-forever charges fit-out "
    "replacement (user cost): 0.70–0.82, under the cash floor at every rate — the old knife-edge framing is "
    "retired. Leased-site rent contingency: ~0.60–0.65 if the three unnamed sites are leased from private "
    "landlords ($11–27M of rent checks not in the $167M bill).",
]
put_label(ws, r, 1, "Reading this tab", bold=True); r += 1
for n in update_notes:
    c = ws.cell(row=r, column=1, value=f"• {n}")
    c.font = label_font
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 30
    r += 1


# =============================================================================
# 6. ALTERNATIVES — food-and-grocery benefit per public dollar
# =============================================================================
ws = wb.create_sheet("Alternatives")
ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 14
ws.column_dimensions["E"].width = 16
ws.column_dimensions["F"].width = 16
ws.column_dimensions["G"].width = 12
ws.column_dimensions["H"].width = 60

put_header(ws, 1, "Alternatives — food-and-grocery benefit per public dollar", span=8)
put_label(ws, 2, 1,
          "The metric is $ of food-and-grocery benefit delivered per $ of public expenditure. "
          "Yellow cells are editable assumptions; leverage = benefit ÷ cost. Each row's title tags "
          "what its numerator is: (access value) / (induced savings) / (cash-like) / (meals delivered) / (price transfer).",
          note=True)
put_label(ws, 3, 1,
          "Caveat: per-household benefit estimates are rough midpoints. Interventions also differ in nutritional quality, "
          "targeting precision, and durability. Edit the yellow cells to see how your own estimates change the ranking.",
          note=True)

r = 5
# Header row
headers = ["#", "Intervention", "10-yr public cost", "HH / kids reached",
           "Avg benefit per HH/yr", "10-yr benefit delivered", "Leverage",
           "Basis for benefit estimate"]
for j, h in enumerate(headers):
    c = ws.cell(row=r, column=1 + j, value=h)
    c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    c.fill = header_fill
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[r].height = 36
r += 1

# ---------------------------------------------------------------
# Interventions:
#   (id, name, cost_formula_annual, reach_formula, benefit_per_hh, basis_note)
# ---------------------------------------------------------------
# 10-yr public cost is expressed directly (not annual × 10) so front-loaded
# programs don't get mis-scaled.
interventions = [
    ("A1",
     "New supermarkets via FRESH-style abatements (access value)",
     f"={A['fresh_per_store']}*93",
     f"={U['low_access_hh']}",
     300,
     "Point estimate ~4.3×; honest range 2–8× depending on the access-benefit "
     "estimate ($150–600/HH/yr bounds; Allcott-Diamond-Dubé 2019 and NYC's FRESH "
     "evaluation support the low end for typical households, the full $300 only "
     "for the low-income core of the low-access population, ~32,000 households). "
     "Catalyzed private capital: "
     "~$5.9M/store (NYCEDC program page: $177M across 30 completed projects)."),
    ("A2",
     "Costco / BJ's membership subsidy for every SNAP household (10 yrs) (induced savings)",
     f"={A['warehouse_fee']}*{A['snap_households']}*10",
     f"={A['snap_households']}*{A['warehouse_util']}",
     600,
     "70% active use is an author assumption — documented membership-shaped "
     "analogs run 7–46%, necessity-benefit take-up 75–88% (SNAP FY22, USDA). "
     "Club prices run 20–35% below typical supermarkets (Consumers' Checkbook "
     "2022; Consumer Reports 2025, percentage wedges — the $800–1,200/yr and "
     "blended ~$600/yr dollar figures are this model's conversion at "
     "$4,000–6,000/yr food-at-home spend; car/transit split is our assumption)."),
    ("A3",
     "Expand school meals to weekends (full public cost) (meals delivered)",
     f"={A['nycha_students']}*{A['meals_per_kid']}*{A['meal_cost']}*10",
     f"={A['nycha_students']}",
     786,
     "Priced at full public cost — federal free rates run $2.46–5.15/meal "
     "(SY2025-26 NSLP/SBP/CACFP); ~1.0× by construction: meals delivered ≈ "
     "dollars spent. (April had priced these at the state top-up $0.35/meal, "
     "which manufactured a 10×; corrected August 2026. NYC summers are already "
     "universal — SFSP + SUN Bucks — so the expansion is weekends.)"),
    ("A4",
     "Bodega produce upgrade (one-time grant + 2-yr subsidy, then monitoring) (access value)",
     f"={A['bodegas']}*5000+{A['bodegas']}*2600",
     f"={A['nyc_households']}",
     50,
     "Reach is the deduplicated citywide ceiling (3,334,088 NYC households, "
     "Census QuickFacts 2020–2024) — nearly every NYC household lives near a "
     "bodega, so per-bodega catchment arithmetic double-counts. ~$50–100/HH/yr "
     "realized produce-access value; $5k one-time infra grant + $2,600/bodega "
     "2-yr produce subsidy. Point ~22× at the $76M natural scope, ~10–20× "
     "normalized to the plan's envelope (depth scales, reach cannot). Was 32.9× "
     "on overlapping catchments — corrected August 2026."),
    ("A5",
     "Expand Health Bucks SNAP matching to $40/day, 200 markets (10 yrs) (cash-like)",
     300000000,
     f"={A['hb_reach']}",
     150,
     "Doubles fresh-produce purchasing for participating SNAP users. "
     "Est. $150/yr of incremental produce per participant."),
    ("A6",
     "NYC SNAP supplement ($20/mo) — SNAP families with kids (10 yrs) (cash-like)",
     f"={A['snap_topup']}*12*{A['snap_households']}*{A['snap_kids_share']}*10",
     f"={A['snap_households']}*{A['snap_kids_share']}",
     240,
     "~40% of SNAP HH have kids. Direct transfer: every public $ = $1 of food "
     "the shopper chose. Leverage is 1.0x by design — this is the honest floor."),
    ("PLAN",
     "Mayor Mamdani's 5 city-owned grocery stores (central case) (price transfer)",
     f"={A['total_capex']}+{A['annual_subsidy']}*10",
     "90000",
     67,
     "~18k shoppers × 5 stores × $1,100/yr spend at the store × ~6% realized discount "
     "(the unit economics can't support a double-digit discount — see Sensitivity tab). "
     "Matches the April piece's published central-case row (90k shoppers, $67/HH/yr, 0.6x); "
     "resynced August 2026 — the workbook previously shipped 75k × $50, which did not "
     "match the published table. Modeled April 2026, before any official discount figure "
     "existed; for the 30% core-basket discount announced July 27, 2026 see the July 2026 "
     "Update tab. Benefit is price savings on groceries the shopper was already buying, "
     "not new nutrition."),
    ("PLAN+",
     "Mayor Mamdani's plan — charitable case (plan assumptions work) (price transfer)",
     f"={A['total_capex']}+{A['annual_subsidy']}*10",
     "100000",
     130,
     "Generous: 20k shoppers × 5 stores, ~11% realized discount on $1,200/yr store spend. "
     "Matches the April piece's published charitable row (100k shoppers, $130/HH/yr, 1.3x); "
     "resynced August 2026 from the previously shipped $120/HH. "
     "Requires sales/sqft + cost structure the plan has not demonstrated it can achieve."),
]

# Cost-cell comments where the cost formula needs explaining (keyed by row id).
cost_comments = {
    "A1": "93 new stores × $1.08M public cost per store (Comptroller Fiscal Note "
          "4-2024) ≈ $100M — the store count is chosen to normalize this row to "
          "a ~$100M ten-year envelope for comparability across rows.",
}

for idx, (num, name, cost_f, reach_f, benefit_hh, note) in enumerate(interventions):
    put_label(ws, r, 1, num, bold=True)
    c = ws.cell(row=r, column=2, value=name)
    c.font = label_font
    c.alignment = Alignment(wrap_text=True, vertical="top")

    # 10-yr public cost (col C) — each row specifies 10-yr total directly
    cost_cell = ws.cell(row=r, column=3)
    cost_cell.value = cost_f
    cost_cell.number_format = FMT_D
    cost_cell.font = label_font
    if num in cost_comments:
        cost_cell.comment = Comment(cost_comments[num], "17A")

    # Reach (col D)
    reach_cell = ws.cell(row=r, column=4)
    if isinstance(reach_f, str) and reach_f.startswith("="):
        reach_cell.value = reach_f
    elif isinstance(reach_f, str):
        reach_cell.value = float(reach_f)
    else:
        reach_cell.value = reach_f
    reach_cell.number_format = FMT_N
    reach_cell.font = label_font

    # Avg benefit/HH/yr (col E) — editable yellow
    put_input(ws, r, 5, benefit_hh, fmt=FMT_D)

    # 10-yr benefit delivered (col F) = reach × benefit × 10
    benefit_total = ws.cell(row=r, column=6)
    benefit_total.value = f"=D{r}*E{r}*10"
    benefit_total.number_format = FMT_D
    benefit_total.font = label_bold

    # Leverage (col G) = benefit ÷ cost
    lev = ws.cell(row=r, column=7)
    lev.value = f"=IFERROR(F{r}/C{r},\"-\")"
    lev.number_format = '0.0"x"'
    lev.font = label_bold
    lev.alignment = Alignment(horizontal="right")

    # Basis note (col H)
    note_cell = ws.cell(row=r, column=8)
    note_cell.value = note
    note_cell.font = note_font
    note_cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Highlight the plan row for contrast
    if num == "PLAN":
        for col in [1, 2, 3, 4, 6, 7, 8]:
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor="FADBD8")
    elif num == "PLAN+":
        for col in [1, 2, 3, 4, 6, 7, 8]:
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor="FDF2E9")

    ws.row_dimensions[r].height = 64
    r += 1

r += 2
put_label(ws, r, 1, "How to read this table", bold=True, note=False); r += 1
read_notes = [
    "Leverage = $ of food-and-grocery benefit delivered to households ÷ $ of public expenditure. Numerators differ in kind — each row's title carries a tag (access value / induced savings / cash-like / meals delivered / price transfer); the update piece's methodology notes carry the full taxonomy.",
    "• 1.0x = a direct transfer (every public dollar becomes a dollar of food benefit). This is the honest floor for any food-access intervention.",
    "• >1.0x = catalytic leverage — public spend unlocks more food benefit than it costs (private capital, existing infrastructure, volume discounts).",
    "• <1.0x = public spend delivers less food benefit than it costs. In the grocery plan's case, it's because most of the spend goes to construction and ongoing subsidy of a retail operation, not to food.",
    "",
    "The five-store plan is the only row below 1.0x even in its charitable case; the central case is 0.6x on this table's reach × spend frame (the April piece's published frame). The July 2026 Update tab prices the announced 30% discount on a throughput frame — 0.69x in today's dollars.",
    "Direct transfers (SNAP supplement, Health Bucks) return ~1x by design and should be considered the baseline any food-access intervention must beat.",
    "Catalytic rows run roughly 2–22x — FRESH-style abatements 2–8x depending on the access-benefit estimate, warehouse memberships ~6.5x in induced savings, the bodega ceiling ~10–22x depending on scope.",
]
for n in read_notes:
    c = ws.cell(row=r, column=1, value=n if not n.startswith("•") else n)
    c.font = label_font if not n.startswith("•") else note_font
    c.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 18
    r += 1


# =============================================================================
# 7. SOURCES
# =============================================================================
ws = wb.create_sheet("Sources")
ws.column_dimensions["A"].width = 50
ws.column_dimensions["B"].width = 80
put_header(ws, 1, "Sources", span=2)

sources = [
    ("Mayor's Office, La Marqueta announcement (April 14, 2026)",
     "First site; $30M capex; operator structure ('contractually required to pass "
     "savings directly to customers'). No discount percentage named."),
    ("Mayor's Office, The Peninsula announcement (May 18, 2026)",
     "Second site: Hunts Point, Bronx; 20,000 sqft; first store to open (end 2027)."),
    ("Mayor's Office, 30% discount announcement (July 27, 2026)",
     "30% off a core basket (all fresh produce, meat, seafood + ~20 categories of "
     "pantry staples, dairy, refrigerated); projected 15% cut to the average total "
     "grocery bill ($90/mo ≈ $1,000/yr per shopper); no income test; prices locked "
     "monthly; NYC Groceries private label; operator RFP due Oct 16, 2026. First "
     "hard discount figure in the plan's history — campaign and spring 2026 "
     "materials promised only 'wholesale prices' / passing on savings."),
    ("DoD commissaries — DeCA patron savings & appropriations (FY2024–26); GAO-17-80",
     "External anchor: ~24–25% patron savings (DeCA self-reported; GAO-17-80 "
     "criticized the methodology) on a $1.53–1.57B/yr appropriation (FY2024–26), "
     "236-store scale."),
    ("N.Y.C. Groceries vision plan — 'A Recipe for Affordability' (July 27, 2026)",
     "The program's operative document: basket = all fresh produce/meats/seafood + "
     "select dairy, shelf-stable, frozen; priced 'on average 30 percent below market'; "
     "fair pricing required on non-basket items; Peninsula respecified at 15,000 sqft; "
     "no dollar figures anywhere. Local copy: research/vision_plan.pdf."),
    ("NYCEDC — N.Y.C. Groceries Operator(s) RFP (July 27, 2026; 44 pp., obtained Aug 2 2026)",
     "Footprints (App. D–E), Core Basket categories (App. C), Affordability Payment "
     "structure (pp. 6–7, 20): annual payments sized to the operating deficit from "
     "discounted Core Basket sales, net of free rent and property tax; city-funded "
     "buildout; Labor Peace Agreement; proposals due Oct 16, 2026; operators "
     "announced early 2027."),
    ("NYC Council Finance Division / Comptroller (Feb–Mar 2026)",
     "$70M entered the capital record in the Feb 2026 Preliminary Commitment Plan, all "
     "planned FY2027, flowing through EDC capital (DSBS). No named budget line yet; "
     "no operating subsidy appears anywhere in the adopted FY2027 budget."),
    ("Implementation ledger (this repo)",
     "gizmos/nyc-public-grocery-math/IMPLEMENTATION_LEDGER_2026-08.md — every "
     "official document and quote, sourced and dated, with conflicts flagged."),
    ("FMI, The Food Retailing Industry Speaks 2026 (rel. July 7, 2026)",
     "2025 weekly sales per sqft of selling area $19.59 ≈ $1,019/yr. FMI 2024 Food "
     "Industry Facts remains the anchor for COGS %, average store size (42k sqft), "
     "and the ~1.6% industry net margin."),
    ("U.S. Census Bureau — QuickFacts, New York City (2020–2024 ACS)",
     "2.48 persons per household; 3,334,088 households."),
    ("Company filings via stockanalysis.com — Kroger net margin",
     "At or below 1.5% in three of the last five fiscal years; latest ~0.7%."),
    ("Company filings via stockanalysis.com — Costco net margin",
     "~2.9% in the last two fiscal years, membership-fee driven."),
    ("USDA Thrifty Food Plan, 2025",
     "~$750/mo for a mother + 2 school-age kids in NYC."),
    ("NY Department of Labor — 2026 minimum + prevailing wage schedule",
     "Base wage, union supplement, benefits loading."),
    ("NYC Comptroller, Good Jobs and the New York City FRESH Program (Fiscal Note 4-2024, Oct 2024)",
     "$29.2M cumulative tax expenditures through FY2023 across 27 tax-subsidized "
     "stores (~$1.08M/store); 30-completed-projects and $177M-private-capital "
     "figures are NYCEDC program-page numbers."),
    ("NYS OTDA — SNAP caseload statistics (data.ny.gov dq6j-8u8z)",
     "2025 monthly avg: ~1,069,000 NYC SNAP households, ~1.77M recipients "
     "(~1.66 persons/household)."),
    ("NY Universal Free School Meals (enacted May 2025, SY2025-26)",
     "~$340M/yr state top-up on top of federal NSLP/SBP; federal free rates run "
     "$2.46–5.15/meal (SY2025-26 NSLP/SBP/CACFP) and carry most of the cost of "
     "new meals."),
    ("NYC Food Policy Center — East Harlem food environment",
     "Retail density, bodega counts, food-access metrics."),
    ("The Real Deal — FRESH analysis",
     "Property-level review of FRESH program results."),
    ("KCUR — Kansas City Sun Fresh closure (2025)",
     "Municipal grocery precedent; ~$18–21M total public cost."),
    ("Action News Jax — Baldwin FL municipal grocery closure",
     "Precedent; closed after 5 years."),
    ("Supermarket News — Chicago drops municipal grocery (2024)",
     "Mayor Johnson backed away from the concept after feasibility study."),
    ("NYC DOHMH — Health Bucks program (distributed with GrowNYC)",
     "Current $10/day cap; match-rate evaluation for fresh produce."),
]

r = 3
for name, desc in sources:
    put_label(ws, r, 1, name, bold=True)
    c = ws.cell(row=r, column=2, value=desc)
    c.font = label_font
    c.alignment = Alignment(wrap_text=True)
    r += 1


# =============================================================================
# Visual polish (D-18 — presentation only, no content)
# =============================================================================
# Gridlines off everywhere (fills and borders carry the structure); brand tab
# colors (Cobalt for the update tab, gold for the README, charcoal elsewhere);
# freeze panes so headers stay put on the data tabs.
TAB_COLORS = {"July 2026 Update": COBALT, "README": "FFBD3D"}
FREEZE_PANES = {
    "Per-Store P&L": "B5",       # rows 1–4 (year header) + label column
    "Five-Store Rollup": "B5",   # rows 1–4 (year header) + label column
    "Sensitivity": "B6",         # rows 1–5 (grid header) + label column
    "July 2026 Update": "A4",    # header + intro note
    "Alternatives": "A6",        # header + column-header row (data starts row 6)
}
for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = TAB_COLORS.get(sheet.title, CHARCOAL)
    if sheet.title in FREEZE_PANES:
        sheet.freeze_panes = FREEZE_PANES[sheet.title]

# With gridlines off, the two dense numeric grids need their own structure:
# thin borders on every populated cell of the Sensitivity grid and the
# Alternatives table body (everything else is carried by fills already).
box = Border(top=thin, bottom=thin, left=thin, right=thin)
for title, min_row in (("Sensitivity", 5), ("Alternatives", 5)):
    sheet = wb[title]
    for row in sheet.iter_rows(min_row=min_row):
        for cell in row:
            if cell.value is not None and not isinstance(cell.value, str) or (
                    isinstance(cell.value, str) and cell.value.strip() and not cell.value.startswith("•")):
                if cell.column <= 8 and not cell.alignment.wrap_text:
                    cell.border = box

# =============================================================================
# Save
# =============================================================================
# Set README as the active tab on open
wb.active = wb.index(wb["README"])
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(f"Wrote {OUTPUT}")
print(f"Size: {OUTPUT.stat().st_size / 1024:.1f} KB")
