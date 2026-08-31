#!/usr/bin/env python3
"""Verification gate — The New Math on NYC's Public Grocery Stores (v3, 2026-08-06).

v3 (economist-review round, D-22/D-23): asset-value/symmetry/real-rate/user-cost blocks added to
section 0 and the workbook checks; the perpetuity knife-edge prose retired to MUST_NOT;
editor's note, sell-the-stores subsection, and self-funded-ceiling pins added.

Rewritten after the August 2026 adversarial pass (research/audit_2026-08/): the model
re-centered on the RFP-consistent frame (13,800 selling sqft x FMI-2025 $1,019/sqft,
18% blend, Census 2.48 persons/HH) with leverage in present value at the model's 4%
municipal rate ("in today's dollars"). Decisions in DECISION_LEDGER.md; canonical values
in CANONICAL_NUMBERS.md. Sections:

  0. Ground truth: every headline number re-derived from the input constants (derived,
     not hard-coded, wherever an identity exists — a changed assumption should scream).
  1. Web dial config == model constants (src/config/groceryMath30.ts), incl. the
     basketDiscount x basketShare == defaultBlend identity.
  2. LibreOffice-evaluated workbook cells == derivations (July 2026 Update tab: grids,
     diagonal, horizon block, same-money block, variants; Alternatives leverage cells;
     Assumptions April-convention pin). Label-indexed, not cell-addressed.
  3. Prose: load-bearing literals present/absent; no self-closing embed tags; word cap;
     semicolon + em-dash budgets. NEW: GroceryDial.tsx copy literals (the dial narrates
     the model; its hard-coded copy is pinned here so it cannot drift silently).
  4. Charts: the grocery-30 data arrays parsed NUMERICALLY from build_charts.py and
     compared to derivations (replaces the old token-presence greps), plus footnote
     literal pins.
  5. Freshness: build_model.py's own source hash must match the fingerprint the build
     stamped into the workbook README tab (catches edit-without-rebuild).

Needs LibreOffice (soffice) on PATH for section 2; fails cleanly (not a traceback) if
missing. Run: python3 gizmos/nyc-public-grocery-math/verify_claims.py
"""
from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import os
import sys
import tempfile
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("FAIL: openpyxl required (pip install openpyxl)")
    sys.exit(1)

ROOT = Path(__file__).resolve().parent.parent.parent
if not (ROOT / "src").exists():  # draft runs from the scratchpad; final lives in gizmos/<slug>/
    _root = os.environ.get("GIZMO_WAREHOUSE_ROOT")
    if not _root:
        sys.exit("Run from the repo checkout, or set GIZMO_WAREHOUSE_ROOT to it.")
    ROOT = Path(_root)
XLSX = ROOT / "public" / "assets" / "nyc-public-grocery-model.xlsx"
CONTENT_TS = ROOT / "src" / "content" / "nyc-public-grocery-new-math.ts"
APRIL_TS = ROOT / "src" / "content" / "nyc-public-grocery-math.ts"
CONFIG_TS = ROOT / "src" / "config" / "groceryMath30.ts"
DIAL_TSX = ROOT / "src" / "components" / "grocery" / "GroceryDial.tsx"
CHARTS_PY = ROOT / "gizmos" / "nyc-public-grocery-math" / "build_charts.py"
MODEL_PY = ROOT / "gizmos" / "nyc-public-grocery-math" / "build_model.py"

FAILURES: list[str] = []
N_CHECKS = 0


def check(name: str, ok: bool, detail: str = "") -> None:
    global N_CHECKS
    N_CHECKS += 1
    if ok:
        print(f"  PASS  {name}")
    else:
        FAILURES.append(f"{name}  {detail}")
        print(f"  FAIL  {name}  {detail}")


def close(a, b, tol=0.005) -> bool:
    return a is not None and b is not None and abs(a - b) <= tol


# =============================================================================
# 0. GROUND TRUTH — the model chain from first principles
# =============================================================================
print("== 0. Ground truth (independent re-derivation) ==")

# Input constants (the ONLY hand-entered numbers in this section; everything
# below derives from them). Change one and every dependent check re-centers.
STORES, AVG_SQFT, YEARS = 5, 13_800, 10
COGS, LABOR, OPEX = 0.72, 0.12, 0.09
COST_RATIO = COGS + LABOR + OPEX          # 0.93
MARGIN = 1 - COST_RATIO                    # 0.07 pre-discount margin
FEE, CAPEX = 400_000, 70_000_000
BASKET_DISC, BASKET_SHARE = 0.30, 0.60
BLEND = BASKET_DISC * BASKET_SHARE         # 0.18
REV, REV_APRIL = 1_019, 500
DISC_APRIL = 0.10
PPH = 2.48
SAVINGS_MO = 90
RATE = 0.04
AF = (1 - (1 + RATE) ** -YEARS) / RATE     # 10-yr annuity factor at 4%
SNAP_TOPUP_MO = 20
CLUB_FEE, CLUB_UTIL = 65, 0.70
FRESH_COST = 1_080_000                     # Comptroller FN 4-2024: $29.2M / 27 stores
MEAL_COST, MEALS_PER_KID = 4.91, 160       # CACFP at-risk supper analog; weekends only


def chain(rev: float, blend: float, sqft: float = AVG_SQFT) -> dict:
    g = sqft * rev
    result = g * (MARGIN - blend) - FEE
    subsidy = -STORES * result
    envelope = CAPEX + YEARS * subsidy
    transfer = STORES * g * blend
    hh = transfer / (12 * SAVINGS_MO)
    env_pv = CAPEX + subsidy * AF
    ben_pv = transfer * AF
    return {
        "gross": g, "result": result, "subsidy": subsidy, "envelope": envelope,
        "transfer": transfer, "hh": hh, "people": hh * PPH,
        "env_pv": env_pv, "ben_pv": ben_pv,
        "lev_pv": ben_pv / env_pv, "lev_nom": transfer * YEARS / envelope,
    }


C = chain(REV, BLEND)                      # announced central
A = chain(REV_APRIL, DISC_APRIL)           # April column at the current frame
Q = chain(REV_APRIL, BLEND)                # quiet fallback
TJ = chain(2_000, BLEND)                   # Trader Joe's-class

check("central gross/store $14,062,200", close(C["gross"], 14_062_200, 1))
check("central result/store -$1,946,842", close(C["result"], -1_946_842, 1))
check("central subsidy $9,734,210/yr", close(C["subsidy"], 9_734_210, 1))
check("central envelope $167,342,100 ('about $167M')", close(C["envelope"], 167_342_100, 1))
check("central transfer $12,655,980/yr ('about $13M')", close(C["transfer"], 12_655_980, 1))
check("central households 11,718.5 ('roughly 12,000')", close(C["hh"], 11_718.5, 0.1))
check("central people 29,061.9 ('about 29,000')", close(C["people"], 29_061.88, 0.1))
check("central PV leverage 0.6892 ('0.69x')", close(C["lev_pv"], 0.689152, 0.0005))
check("april column envelope ~$100.35M (nominal, = published '$100M')",
      close(A["envelope"], 100_350_000, 1))
check("april column PV leverage 0.2957 ('0.30x')", close(A["lev_pv"], 0.295747, 0.0005))
check("'more than doubles': PV ratio 2.33", close(C["lev_pv"] / A["lev_pv"], 2.3302, 0.001))
check("quiet fallback envelope $127.95M ('about $128M')", close(Q["envelope"], 127_950_000, 1))
check("quiet fallback PV leverage 0.4305 ('0.43x')", close(Q["lev_pv"], 0.430492, 0.0005))
check("TJ's-class stays under the floor in PV", TJ["lev_pv"] < 1.0,
      f"lev_pv={TJ['lev_pv']:.4f}")

# PV crossover: transfer*AF = CAPEX + subsidy*AF  =>  blend cancels entirely.
G_STAR = (CAPEX + STORES * FEE * AF) / (STORES * MARGIN * AF)
check("PV crossover $2,200.9/sqft ('about $2,200'), blend-independent",
      close(G_STAR / AVG_SQFT, 2_200.9, 0.5))
for b, expect in ((BLEND, 18_704_861), (0.30, 36_928_346)):
    sub_x = STORES * (G_STAR * (b - MARGIN) + FEE)
    check(f"crossover subsidy @{b:.0%} blend ${expect/1e6:.1f}M/yr", close(sub_x, expect, 5))

check("asymptotic ceiling @18% = blend/(blend-margin) = 1.636",
      close(BLEND / (BLEND - MARGIN), 1.636364, 0.0005))
check("asymptotic ceiling @30% = 1.304", close(0.30 / (0.30 - MARGIN), 1.304348, 0.0005))

# Horizon block: the flank the piece closes with one clause.
check("operating leverage (transfer/subsidy) 1.300", close(C["transfer"] / C["subsidy"], 1.3002, 0.001))
AF30 = (1 - (1 + RATE) ** -30) / RATE
check("30-yr PV leverage 0.918",
      close(C["transfer"] * AF30 / (CAPEX + C["subsidy"] * AF30), 0.9183, 0.001))
check("perpetuity PV leverage ~1.01 ('roughly match... on day one')",
      close((C["transfer"] / RATE) / (CAPEX + C["subsidy"] / RATE), 1.009715, 0.001))
STORE_YEARS_PHASED = 44                    # opens 1/2/2 across years 1-3, window ends yr 10
check("phased envelope $155.66M",
      close(CAPEX + STORE_YEARS_PHASED * -C["result"], 155_661_048, 1))

# Same-money block (nominal annual pool, per the table's convention).
POOL = C["envelope"] / YEARS
check("pool $16,734,210/yr", close(POOL, 16_734_210, 1))
check("equal-depth cash 15,494.6 HH ('~15,500')", close(POOL / (12 * SAVINGS_MO), 15_494.6, 0.1))
check("cash people 38,426.7 ('38,000')", close(POOL / (12 * SAVINGS_MO) * PPH, 38_426.7, 0.5))
SNAP_HH = POOL / (SNAP_TOPUP_MO * 12)
check("SNAP-style 69,725.9 HH ('~70,000')", close(SNAP_HH, 69_725.9, 0.5))
check("SNAP-style people 172,920 ('173,000')", close(SNAP_HH * PPH, 172_920.2, 1))
CLUB_FUNDED = POOL / CLUB_FEE
check("warehouse funded 257,449 ('257,000')", close(CLUB_FUNDED, 257_449.4, 1))
check("warehouse active 180,215 ('~180,000')", close(CLUB_FUNDED * CLUB_UTIL, 180_214.6, 1))
check("warehouse people 446,932 ('447,000')", close(CLUB_FUNDED * CLUB_UTIL * PPH, 446_932.1, 2))
check("reach ratio: SNAP-style/stores 5.95 ('six times')", close(SNAP_HH / C["hh"], 5.9501, 0.005))
check("reach ratio: warehouse-active/stores 15.38 ('fifteen times')",
      close(CLUB_FUNDED * CLUB_UTIL / C["hh"], 15.3786, 0.01))
check("school meals: $786/kid/yr full public cost ('≈$790')",
      close(MEAL_COST * MEALS_PER_KID, 785.60, 0.01))
check("school meals: ~21,300 kids at the pool ('~21,000')",
      close(POOL / (MEAL_COST * MEALS_PER_KID), 21_301.2, 1))
check("FRESH: ~155 stores over ten years ('~150')",
      close(C["envelope"] / FRESH_COST, 154.95, 0.1))
check("implied basket spend $3,600/yr", close(12 * SAVINGS_MO / BASKET_DISC, 3_600, 0.01))
check("30 cents of every public dollar to the wrapping (1 - 0.69)",
      close(1 - C["lev_pv"], 0.3108, 0.001))

# Footprint variants (the '$156M and $173M' span) + full-discount bounding case.
check("footprint variant 11,800 -> $156.13M", close(chain(REV, BLEND, 11_800)["envelope"], 156_133_100, 1))
check("footprint variant 14,800 -> $172.95M", close(chain(REV, BLEND, 14_800)["envelope"], 172_946_600, 1))
check("full-discount bounding case $196.56M ('closer to $197M')",
      close(CAPEX + YEARS * C["transfer"], 196_559_800, 1))

# Economist-review round (Aug 2026, D-22/D-23): asset value, symmetry, real rates, user cost.
TV_LOW, TV_HIGH, TV_GEN = 3e6, 12e6, 30e6
RENT_PSF, G_INFL, FITOUT_LIFE = 40, 0.0226, 12.5
B10 = (1 + RATE) ** -YEARS


def lev_tv(tv: float) -> float:
    return C["ben_pv"] / (C["env_pv"] - tv * B10)


check("sale credited, researched low $3M -> 0.6987 ('0.70')", close(lev_tv(TV_LOW), 0.698658, 0.0005))
check("sale credited, researched high $12M -> 0.7288 ('0.73')", close(lev_tv(TV_HIGH), 0.728818, 0.0005))
check("sale credited, generous $30M -> 0.7977 ('0.80')", close(lev_tv(TV_GEN), 0.797687, 0.0005))
TV_FLOOR = (C["env_pv"] - C["ben_pv"]) / B10
check("sale needed for the 1.0x floor $68.54M ('$68.5M')", close(TV_FLOOR, 68_538_016, 5))
PERP = (C["transfer"] / RATE) / (CAPEX + C["subsidy"] / RATE)
check("IDENTITY: full-$70M sale credit == perpetuity ratio exactly", close(lev_tv(CAPEX), PERP, 1e-9))
RENT_5 = STORES * AVG_SQFT * RENT_PSF
check("implied market rent, five stores $2.76M/yr", close(RENT_5, 2_760_000, 1))
check("full-balance-sheet (rent charged) low 0.6063 ('0.61')",
      close(C["ben_pv"] / (C["env_pv"] + RENT_5 * AF - TV_LOW * B10), 0.606283, 0.0005))
check("full-balance-sheet (rent charged) high 0.6289 ('0.63')",
      close(C["ben_pv"] / (C["env_pv"] + RENT_5 * AF - TV_HIGH * B10), 0.628866, 0.0005))
LEASED = 3 * 15_000 * RENT_PSF
check("leased-site rent $1.8M/yr at $40/sqft -> 0.6276",
      close(LEASED, 1_800_000, 1)
      and close(C["ben_pv"] / (C["env_pv"] + LEASED * AF), 0.627634, 0.0005))
for psf, expect in ((25, 0.6494), (60, 0.6008)):  # the '$25-60 band -> ~0.60-0.65' bracket
    check(f"leased-site band @${psf}/sqft -> {expect}",
          close(C["ben_pv"] / (C["env_pv"] + 3 * 15_000 * psf * AF), expect, 0.001))
R_REAL = (1 + RATE) / (1 + G_INFL) - 1
AF_REAL = (1 - (1 + R_REAL) ** -YEARS) / R_REAL
check("real-rate ten-year read 0.7271 ('0.72-0.73 rather than 0.689')",
      close(C["transfer"] * AF_REAL / (CAPEX + C["subsidy"] * AF_REAL), 0.727110, 0.0005))
check("no-replacement indexed perpetuity 1.156 ('about 1.16')",
      close(C["transfer"] / (CAPEX * (RATE - G_INFL) + C["subsidy"]), 1.155564, 0.001))
USER_COST = CAPEX * (R_REAL + 1 / FITOUT_LIFE)
check("user cost K(r+1/life) $6.79M/yr", close(USER_COST, 6_791_082, 2))
check("run forever with replacement 0.766 central",
      close(C["transfer"] / (C["subsidy"] + USER_COST), 0.765855, 0.0005))
for g, life, expect in ((0.025, 15, 0.8205), (0.02, 10, 0.6989)):  # 'between 0.70 and 0.82'
    rr = (1 + RATE) / (1 + g) - 1
    uc = CAPEX * (rr + 1 / life)
    check(f"forever band corner (g={g:.1%}, life={life}) -> {expect}",
          close(C["transfer"] / (C["subsidy"] + uc), expect, 0.001))

# =============================================================================
# 1. WEB DIAL CONFIG == MODEL
# =============================================================================
print("== 1. Dial config (src/config/groceryMath30.ts) ==")
cfg = CONFIG_TS.read_text()

def cfg_num(key: str):
    m = re.search(rf"^\s*{key}:\s*([0-9._eE]+)", cfg, re.M)
    return float(m.group(1).replace("_", "")) if m else None

for key, want in [
    ("stores", STORES), ("avgSqft", AVG_SQFT), ("cogsPct", COGS), ("laborPct", LABOR),
    ("opexPct", OPEX), ("operatorFee", FEE), ("capexTotal", CAPEX), ("years", YEARS),
    ("discountRate", RATE), ("basketDiscount", BASKET_DISC), ("basketShare", BASKET_SHARE),
    ("promisedSavingsMonthly", SAVINGS_MO), ("defaultRevPerSqft", REV),
    ("defaultBlend", BLEND), ("defaultPersonsPerHH", PPH),
    ("aprilRevPerSqft", REV_APRIL), ("aprilDiscount", DISC_APRIL),
]:
    check(f"config {key} == {want}", close(cfg_num(key), want, 1e-9))
check("identity: basketDiscount x basketShare == defaultBlend",
      close(cfg_num("basketDiscount") * cfg_num("basketShare"), cfg_num("defaultBlend"), 1e-9))

# =============================================================================
# 2. WORKBOOK (LibreOffice-evaluated) == DERIVATIONS
# =============================================================================
print("== 2. Workbook (LibreOffice-evaluated) ==")
if shutil.which("soffice") is None:
    check("LibreOffice (soffice) on PATH", False, "install LibreOffice or add soffice to PATH")
else:
    with tempfile.TemporaryDirectory() as td:
        try:
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", td, str(XLSX)],
                check=True, capture_output=True, timeout=240)
            evaluated = Path(td) / XLSX.name
            wb = openpyxl.load_workbook(evaluated, data_only=True)
        except Exception as e:  # clean FAIL, not a traceback
            check("workbook evaluates under LibreOffice", False, repr(e))
            wb = None

    if wb is not None:
        july = wb["July 2026 Update"]

        def cells(ws, label_prefix: str, col: str) -> list:
            """Values in `col` for every row whose col-A label starts with the prefix."""
            out = []
            for row in ws.iter_rows(min_col=1, max_col=1):
                c = row[0]
                if isinstance(c.value, str) and c.value.strip().startswith(label_prefix):
                    out.append(ws[f"{col}{c.row}"].value)
            return out

        def one(ws, label_prefix, col, which=0):
            vals = cells(ws, label_prefix, col)
            return vals[which] if len(vals) > which else None

        # Grid A (quiet $500) = first occurrence; Grid B (busy) = second.
        check("July grid A envelope @April(10%) == $100.35M",
              close(one(july, "Ten-year public envelope (capex", "B", 0), A["envelope"], 1))
        check("July grid A today's-dollars leverage @April == 0.2957",
              close(one(july, "Food benefit per public dollar", "B", 0), A["lev_pv"], 0.0005))
        check("July grid B result/store @central == -$1,946,842",
              close(one(july, "Operating result per store", "D", 1), C["result"], 1))
        check("July grid B envelope @central == $167,342,100",
              close(one(july, "Ten-year public envelope (capex", "D", 1), C["envelope"], 1))
        check("July grid B households @central == 11,718.5",
              close(one(july, "Households receiving the full promised savings", "D", 1), C["hh"], 0.1))
        check("July grid B New Yorkers @central == 29,061.9",
              close(one(july, "New Yorkers (households", "D", 0), C["people"], 0.1))
        check("July grid B today's-dollars leverage @central == 0.6892",
              close(one(july, "Food benefit per public dollar", "D", 1), C["lev_pv"], 0.0005))
        check("July annuity factor == 8.1109",
              close(one(july, "Annuity factor", "B"), float(AF), 1e-6))
        check("July footprint variant 11,800 == $156.13M",
              close(one(july, "— envelope if the footprint reads 11,800", "B"),
                    chain(REV, BLEND, 11_800)["envelope"], 1))
        check("July footprint variant 14,800 == $172.95M",
              close(one(july, "— envelope if the footprint reads 14,800", "B"),
                    chain(REV, BLEND, 14_800)["envelope"], 1))
        check("July bounding case == $196.56M",
              close(one(july, "Bounding case: city pays the full discount", "B"),
                    CAPEX + YEARS * C["transfer"], 1))
        check("July diagonal nominal envelope C == $167.34M, growth +66.8% vs $100.35M",
              close(one(july, "Ten-year public envelope (nominal", "C"), C["envelope"], 1)
              and close(one(july, "Ten-year public envelope (nominal", "D"),
                        C["envelope"] / A["envelope"] - 1, 0.0005))
        check("July diagonal today's-dollars leverage C == 0.6892",
              close(one(july, "Food benefit per public dollar", "C", 2), C["lev_pv"], 0.0005))
        check("July diagonal New Yorkers == 29,061.9",
              close(one(july, "New Yorkers getting the full deal", "C"), C["people"], 0.1))
        check("July same-money pool == $16,734,210",
              close(one(july, "Annual pool", "B"), float(POOL), 1))
        check("July cash-at-depth == 15,494.6 / 38,426.7 people",
              close(one(july, "Cash at the stores' own depth", "B"), POOL / (12 * SAVINGS_MO), 0.1)
              and close(one(july, "Cash at the stores' own depth", "C"), POOL / (12 * SAVINGS_MO) * PPH, 0.5))
        check("July SNAP-style == 69,725.9 / 172,920 people",
              close(one(july, "SNAP-style supplement", "B"), SNAP_HH, 0.1)
              and close(one(july, "SNAP-style supplement", "C"), SNAP_HH * PPH, 0.5))
        check("July warehouse funded == 257,449",
              close(one(july, "Warehouse-club memberships funded", "B"), CLUB_FUNDED, 0.5))
        check("July warehouse active == 180,215 / 446,932 people",
              close(one(july, "— actively used", "B"), CLUB_FUNDED * CLUB_UTIL, 0.5)
              and close(one(july, "— actively used", "C"), CLUB_FUNDED * CLUB_UTIL * PPH, 1))
        check("July FRESH stores == 154.9 (÷ $1.08M)",
              close(one(july, "FRESH-style abatements: new supermarkets", "B"),
                    C["envelope"] / FRESH_COST, 0.05))
        check("July half-mile low-access households == 143,000 (USDA all incomes, D-17a)",
              close(one(july, "— households beyond a half-mile walk", "B"), 143_000, 1))
        check("July bodega ceiling == 3,334,088 (Census 2020-2024)",
              close(one(july, "Bodega-upgrade route: reach ceiling", "B"), 3_334_088, 1))
        # Horizon block (D1)
        check("July operating leverage == 1.300",
              close(one(july, "Operating leverage (transfer", "B"), C["transfer"] / C["subsidy"], 0.001))
        check("July envelope in today's dollars == $148.95M",
              close(one(july, "Envelope in today's dollars", "B"), float(C["env_pv"]), 2))
        check("July phased-calendar variant == $155.66M",
              close(one(july, "Phased-calendar variant", "B"),
                    CAPEX + STORE_YEARS_PHASED * -C["result"], 1))
        check("July 30-year leverage == 0.918",
              close(one(july, "30-year leverage", "B"),
                    C["transfer"] * AF30 / (CAPEX + C["subsidy"] * AF30), 0.001))
        check("July perpetuity == 1.0097",
              close(one(july, "Run forever (perpetuity", "B"),
                    (C["transfer"] / RATE) / (CAPEX + C["subsidy"] / RATE), 0.001))
        # Asset value & symmetry block (economist-review round, D-22/D-23)
        check("July sale credited low == 0.6987",
              close(one(july, "Leverage with a year-ten sale credited — researched low", "B"),
                    lev_tv(TV_LOW), 0.0005))
        check("July sale credited high == 0.7288",
              close(one(july, "Leverage with a year-ten sale credited — researched high", "B"),
                    lev_tv(TV_HIGH), 0.0005))
        check("July generous $30M == 0.7977",
              close(one(july, "Leverage if the sale recovers $30M", "B"), lev_tv(TV_GEN), 0.0005))
        check("July sale-needed-for-floor == $68.54M",
              close(one(july, "Sale value needed to reach the 1.0x cash floor", "B"), TV_FLOOR, 5))
        check("July rent $40/sqft, five-store rent $2.76M",
              close(one(july, "Market rent, blended", "B"), RENT_PSF, 1e-9)
              and close(one(july, "Implied market rent, five stores", "B"), RENT_5, 1))
        check("July full-balance-sheet low/high == 0.6063/0.6289",
              close(one(july, "Full-balance-sheet leverage — rent charged, low", "B"),
                    C["ben_pv"] / (C["env_pv"] + RENT_5 * AF - TV_LOW * B10), 0.0005)
              and close(one(july, "Full-balance-sheet leverage — rent charged, high", "B"),
                        C["ben_pv"] / (C["env_pv"] + RENT_5 * AF - TV_HIGH * B10), 0.0005))
        check("July leased-rent $1.8M -> 0.6276",
              close(one(july, "Leased-site contingency", "B"), LEASED, 1)
              and close(one(july, "Leverage if those rent checks join the bill", "B"),
                        C["ben_pv"] / (C["env_pv"] + LEASED * AF), 0.0005))
        check("July real-rate ten-year read == 0.7271",
              close(one(july, "Ten-year leverage, indexed flows", "B"),
                    C["transfer"] * AF_REAL / (CAPEX + C["subsidy"] * AF_REAL), 0.0005))
        check("July naive indexed perpetuity == 1.156",
              close(one(july, "No-replacement perpetuity, indexed flows", "B"),
                    C["transfer"] / (CAPEX * (RATE - G_INFL) + C["subsidy"]), 0.001))
        check("July user cost == $6.79M, forever w/ replacement == 0.766",
              close(one(july, "User cost of capital", "B"), USER_COST, 2)
              and close(one(july, "Run forever with replacement charged", "B"),
                        C["transfer"] / (C["subsidy"] + USER_COST), 0.0005))
        # April convention pin + Alternatives leverage cells
        assumptions = wb["Assumptions"]
        check("Assumptions discount still 0.10 (April tables intact)",
              close(one(assumptions, "Discount to shoppers vs. market price", "B"), 0.10, 1e-9))
        alts = wb["Alternatives"]
        def alt_lev(title_frag: str):
            for row in alts.iter_rows(min_col=2, max_col=2):
                c = row[0]
                if isinstance(c.value, str) and title_frag in c.value:
                    return alts[f"G{c.row}"].value
            return None
        check("Alternatives FRESH point ~4.27x (range 2-8 in notes)",
              close(alt_lev("FRESH-style abatements"), 4.2712, 0.01))
        check("Alternatives Costco 6.46x", close(alt_lev("Costco / BJ"), 6.4615, 0.01))
        check("Alternatives school meals ~1.0x", close(alt_lev("school meals"), 786 / 785.6, 0.01))
        check("Alternatives bodega ~21.9x natural-scope point",
              close(alt_lev("Bodega produce upgrade"), 21.935, 0.05))
        check("Alternatives PLAN 0.60x (April frame, unchanged)",
              close(alt_lev("central case"), 0.603, 0.005))
        check("Alternatives PLAN+ 1.30x (April frame, unchanged)",
              close(alt_lev("charitable case"), 1.30, 0.005))

# =============================================================================
# 3. PROSE + DIAL COPY
# =============================================================================
print("== 3. Prose ==")
prose = CONTENT_TS.read_text()

MUST_CONTAIN = [
    # headline chain
    "$167M", "roughly 12,000 households", "about 29,000 people", "reaches 38,000",
    "between 70,000 and 180,000", "fiscal 2027",
    "173,000 and 447,000",
    "up from my April model's 0.30&times;",
    "69 cents of food benefit delivered per public dollar spent",
    "a dollar out of the treasury arriving as a dollar off a receipt",
    "Past the four points the stores' margin funds",
    "about $2,200 per square foot",
    "about thirty cents of every public dollar",
    "Even permanent stores stay under it once replacing",
    "two-thirds more than the $100M",
    # Economist-review round (D-22/D-23)
    "made two fair points about this analysis",       # editor's note (Joe's 8/7 de-AI pass)
    "these stores can fund about four",               # Who pays: margin-vs-transfer principle
    "straight cash transfer from taxpayers to the shopper at the register",
    "an 11-cent loss on every dollar sold",           # The new bill: loss derived in place
    "And the buildings do not rescue the math",       # body sale beat (elevated 8/7)
    "sell them for what it paid",
    "#### Suppose the city sells the stores",
    "#suppose-the-city-sells-the-stores",             # nav link
    "an identity, not a coincidence",
    "how deep a discount could the stores fund with no subsidy at all",
    "centered near seven",
    "one-sided test",
    "pay all rent and property taxes (as applicable)",
    "between 0.70 and 0.82",
    "0.72–0.73 rather than 0.689",
    "less favorable to my own conclusion",
    "between $156M and $173M", "closer to $197M",
    "about $128M", "0.43&times;", "1.64&times;",
    "several times under the best shovel-ready alternatives",
    # decisions
    "## The City Sizes the Discount",
    "with no income test",
    "The \"on average\" comes from the vision plan",
    "2.48 New Yorkers per household",
    "wrapped in a grocery store",
    "totally favoring generosity over breadth",
    "excessive bulk purchases",             # Parting concern, R2 wording
    "voluntary free savings card",
    "leaves the operating line unstated",   # D-11 'unstated', not 'silent'
    "This model estimates the bill.",
    "60% of sales",                          # basket-share disclosure
    "40&ndash;80% bracket",
    "$1,019",
    "143,000 half-mile low-access households",
    "roughly 32,000 households",
    "beyond a half-mile walk",               # D-17/D-17a analytical decision
    "(price transfer)", "(access value)", "(induced savings)",  # numerator tags
    "methodology notes", "(#methodology)",
    "## Parting concern",
]
for lit in MUST_CONTAIN:
    check(f"prose contains {lit!r}", lit.lower() in prose.lower())

MUST_NOT = [
    "$153M", "0.67&times;", "0.67\u00d7", "24,000 New Yorkers", "seven to seventeen",
    "$2,180", "2.55", "order of magnitude under",
    "vision plan bars resale", "mechanism unspecified",
    "are silent", "stapled", "not both, and at",
    "Units:", "The first hard number\n",     # old header (allow the in-text phrase)
    "$16&ndash;34M", "$16\u201334M", "up to ~355,000", "up to ~355",       # D-17: no 'up to' hedge
    "footprint readings", "Footprint readings",
    # Economist-review round: the retired knife-edge story must not resurface (D-22)
    "roughly match what plain cash does on day one",
    "The break-even is 4.17%",
    "sign not determinable",
]
for lit in MUST_NOT:
    check(f"prose lacks {lit!r}", lit not in prose)

check("no self-closing embed tags",
      not re.findall(r"<[a-z][a-z0-9-]*\s*/>", prose))

# A heading directly under a closing HTML tag (no blank line) renders as
# literal "## ..." text — react-markdown absorbs it into the raw-HTML block.
# Bit us at </details>\n## Dial it yourself (caught live 2026-08-03).
check("every markdown heading is preceded by a blank line",
      not re.findall(r"[^\n]\n#{1,6} ", prose))

# Word cap. Split on '<details' (id attribute now present); body prose only.
# split at the methodology block specifically: the table's "Assumptions and
# constraints" expander sits mid-body and must NOT exempt the sections after it
body = prose.split('<details id="methodology"')[0]
lines = [ln for ln in body.split("\n")
         if ln.strip() and ln.lstrip()[0] not in "|!<*#"]
text = "\n".join(lines).replace("&mdash;", "—")
text = re.sub(r"&[a-z]+;", "·", text)
n_words = len(text.split())
# Cap raised 1,050 -> 1,400 on 2026-08-03: the adjudicated pass added required
# content (blend disclosure, today's-dollars frame, RFP mechanism specifics,
# corrections note), Joe's dek paragraph moved into the body, and Joe's standing
# direction was clarity over compression (Emma review). The cap exists to stop
# drift; it is a governor, not a target. Body measured ~1,309 at the raise.
# Raised again 1,400 -> 1,500 on 2026-08-07 (D-24): Joe elevated the
# margin-vs-transfer principle into "Who pays", the in-place loss derivation
# into "The new bill", and the sale beat into the body. Body ~1,470 at the raise.
check(f"prose body under 1,500 words ({n_words})", n_words < 1_500)
n_semi = text.count(";")
check(f"semicolons <= 3 ({n_semi})", n_semi <= 3)
n_dash = text.count("—")
check(f"em-dash density ({n_dash} in {n_words})", n_dash <= max(3, n_words // 150))


print("== 3c. April piece (corrected in place per D-13) ==")
april = APRIL_TS.read_text()
for lit in [
    "(April 14, 2026)",
    "corrected three of my April assumptions",
    "~143k half-mile low-access HH (~93 new stores)",
    "$30&ndash;60 realized",
    "~$1&ndash;2B", "**~10&ndash;20&times;**", "**2&ndash;8&times;**", "**~1.0&times;**",
    "low-income core of the low-access population (~32,000 households)",
    "Fiscal Note 4-2024",
    "roughly $500&ndash;$560k/year",
    "3&ndash;6&times;",
]:
    check(f"april contains {lit!r}", lit in april)
for lit in [
    "October 2025", "up to ~143k", "genuinely low-access HH",
    "$500&ndash;$700k", "clear 10&ndash;30&times;", "~$3.3B", "33&times;",
    "~5M HH", "~1M HH (~67 new stores)", "~140k kids",
]:
    check(f"april lacks {lit!r}", lit not in april)
check("april: every markdown heading is preceded by a blank line",
      not re.findall(r"[^\n]\n#{1,6} ", april))

print("== 3b. Dial copy (GroceryDial.tsx) ==")
dial = DIAL_TSX.read_text()
for lit in [
    "about $2,200", "93¢", "1.64×", "1.30×", "0.70–0.82×", "0.69×",
    "13,800 avg", "$70M", "$2M a year", "in today's dollars",
    "present value at a 4% municipal borrowing rate",
    "past any real store", "quieter", "industry avg (FMI 2025)",
    "replacing the fit-out as it wears",
]:
    check(f"dial copy contains {lit!r}", lit in dial)
for lit in ['label: "uncharted"', "$2,180", "2.55", "11,800", "$965", "Infinity×", "-2.50",
            "~1.01×"]:
    check(f"dial copy lacks {lit!r}", lit not in dial)

# =============================================================================
# 4. CHARTS — numeric array comparison (not token grep)
# =============================================================================
print("== 4. Charts (numeric) ==")
charts = CHARTS_PY.read_text()

def chart_floats(varname: str, block: str) -> list[float]:
    """Pull the numeric second element from ("label", value, ...) rows."""
    m = re.search(rf"{varname}\s*=\s*\[(.*?)\n\]", block, re.S)
    if not m:
        return []
    return [float(x) for x in re.findall(r",\s*(-?\d+(?:\.\d+)?)\s*,", m.group(1))]

lev_vals = chart_floats("lev_rows", charts)
check("leverage chart: 8 bars", len(lev_vals) == 8, f"got {len(lev_vals)}")
if len(lev_vals) == 8:
    check("leverage bar april 0.30 == derived", close(lev_vals[0], round(A["lev_pv"], 2), 0.005))
    check("leverage bar announced 0.69 == derived", close(lev_vals[1], round(C["lev_pv"], 2), 0.005))
    check("leverage bars cash/SNAP/school ~1.0", all(close(v, 1.0, 0.05) for v in lev_vals[2:5]))
    check("leverage bar FRESH midpoint 5.0 (range 2-8)", close(lev_vals[5], 5.0, 0.01))
    check("leverage bar warehouse 6.5", close(lev_vals[6], 6.46, 0.06))
    check("leverage bar bodega midpoint 15.0 (range 10-20)", close(lev_vals[7], 15.0, 0.01))

m = re.search(r"subsidy\s*=\s*\[([\d.]+),\s*([\d.]+)\]", charts)
check("envelope chart subsidy components [30.35, 97.34]",
      bool(m) and close(float(m.group(1)), A["subsidy"] * YEARS / 1e6, 0.01)
      and close(float(m.group(2)), C["subsidy"] * YEARS / 1e6, 0.01))
check("envelope chart label '+67%'", "(+67%)" in charts)
check("envelope chart labels operating subsidy to stores", "operating subsidy" in charts)

reach_vals = chart_floats("reach_rows", charts)
check("reach chart: 5 bars", len(reach_vals) == 5, f"got {len(reach_vals)}")
if len(reach_vals) == 5:
    expect = [C["people"] / 1e3, POOL / 1080 * PPH / 1e3, SNAP_HH * PPH / 1e3,
              CLUB_FUNDED * CLUB_UTIL * PPH / 1e3, 355.0]
    for got, want, label in zip(reach_vals, expect,
                                ["stores", "cash", "SNAP-style", "warehouse", "low-access"]):
        check(f"reach bar {label} {want:.1f}k", close(got, round(want, 1), 0.15))
check("reach footnote owns the half-mile decision", "half-mile distance test" in charts)
check("reach footnote names the 32,000-household low-income core", "32,000 households" in charts)
check("no 'up to' hedge on the low-access bar", "up to ~355" not in charts)
check("charts carry no stale $153M/0.67/24,000", all(
    t not in charts for t in ["152.6", "0.67", "24.2", "419.1", "$153M"]))

# =============================================================================
# 5. FRESHNESS — workbook must be built from the current build_model.py
# =============================================================================
print("== 5. Freshness ==")
try:
    wb_raw = openpyxl.load_workbook(XLSX, data_only=False)
    readme = wb_raw["README"]
    stamped = None
    for row in readme.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "sha256 of build_model.py" in c.value:
                m = re.search(r"\b([0-9a-f]{64})\b", c.value)
                stamped = m.group(1) if m else None
    live = hashlib.sha256(MODEL_PY.read_bytes()).hexdigest()
    check("workbook built from the CURRENT build_model.py (fingerprint match)",
          stamped == live,
          "rebuild: python3 gizmos/nyc-public-grocery-math/build_model.py"
          if stamped else "no fingerprint cell found in README tab")
except Exception as e:
    check("freshness fingerprint readable", False, repr(e))


try:
    manifest_path = MODEL_PY.parent / "charts-manifest.json"
    import json as _json
    manifest = _json.loads(manifest_path.read_text())
    stale = []
    for name, want in manifest.items():
        fp = ROOT / "public" / "assets" / name
        got = hashlib.sha256(fp.read_bytes()).hexdigest()
        if got != want:
            stale.append(name)
    check("all five chart PNGs match the build manifest (rebuilt from current script)",
          not stale, f"stale: {stale} — rerun build_charts.py")
except FileNotFoundError as e:
    check("charts manifest present", False, repr(e))

# =============================================================================
print()
if FAILURES:
    print(f"{len(FAILURES)}/{N_CHECKS} CHECKS FAILED:")
    for f in FAILURES:
        print(f"  - {f}")
    sys.exit(1)
print(f"ALL {N_CHECKS} CHECKS PASSED")
